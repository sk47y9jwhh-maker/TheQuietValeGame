from pathlib import Path
from copy import deepcopy
import json,re
from docx import Document
from docx.shared import Mm
from reportlab.pdfgen.canvas import Canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import pypdfium2 as pdfium
R=Path(__file__).resolve().parents[2];W=R/'working/rules-v6_8';A=W/'assets';T=R/'output/artist-tiles-complete/tiles-web';OLD=R/'working/rules-v6_7/assets'
pdfmetrics.registerFont(TTFont('Figure','/System/Library/Fonts/Supplemental/GillSans.ttc',subfontIndex=0))
def figure(name,w,h,items,labels):
 c=Canvas(str(A/(name+'.pdf')),pagesize=(w,h));c.setFillColorRGB(.956,.941,.898);c.rect(0,0,w,h,fill=1,stroke=0)
 for path,x,y,iw,ih in items:c.drawImage(str(path),x,y,iw,ih,preserveAspectRatio=True,anchor='c',mask='auto')
 c.setFillColorRGB(.14,.19,.16);c.setFont('Figure',11)
 for x,y,text in labels:c.drawCentredString(x,y,text)
 c.save();p=pdfium.PdfDocument(str(A/(name+'.pdf')));p[0].render(scale=3).to_pil().save(A/(name+'.png'));p.close()
figure('component-families',660,320,[(T/'c01_lumber_yard-basic.png',10,194,95,82),(T/'c01_lumber_yard-upgraded.png',110,194,95,82),(T/'special-back-23.png',230,194,95,82),(T/'special_the_root_weavers_respite.png',330,194,95,82),(T/'golden_tile_the_golden_charter.png',495,187,115,100)]+[(OLD/(name+'.png'),i*108+25,17,62,108) for i,name in enumerate(['boon','burden','arrival','golden','target-5','target-25'])],[(107,296,'CORE · BASIC / UPGRADED'),(327,296,'SPECIAL · BACK / FACE'),(552,296,'GOLDEN TILE'),(55,145,'BOON'),(163,145,'BURDEN'),(271,145,'ARRIVAL'),(379,145,'GOLDEN BOON'),(487,145,'TARGET'),(595,145,'OPEN TARGET')])
roles=['Vanguard','Knight','Sentinel','Ranger','Warden','Quartermaster'];items=[];labels=[]
for i,role in enumerate(roles):
 p=R/f'tmp/rules-v6_7-qa/The_Quiet_Vale_A6_{role}_Steward_Aid_v6_7/page-1.png';items.append((p,i*110+10,22,90,127));labels.append((i*110+55,5,role))
figure('steward-guides',660,154,items,labels)
# Preserve exact story text. Editorial changes below affect explanations only.
src=R/'output/docx/rulebook-v6_7/The_Quiet_Vale_Main_Rulebook_v6_7_Card_Guide.docx';d=Document(src);edits=[]
def find(prefix):return next(p for p in d.paragraphs if p.text.startswith(prefix))
def replace(prefix,text):
 p=find(prefix);old=p.text;props=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None;p.clear();r=p.add_run(text)
 if props is not None:r._r.insert(0,props)
 edits.append({'before':old,'after':text});return p
def split(prefix,texts):
 p=replace(prefix,texts[0]);after=p._p
 for text in texts[1:]:
  q=d.add_paragraph(text,style=p.style);after.addnext(q._p);after=q._p
 edits[-1]['after']='\n\n'.join(texts)
replace('Together, build Population','Build Population and earn Renown together over 12 rounds. Complete Steward Objectives and Arrivals, resolve Burdens and protect the settlement from Strain. After Round 12, calculate one shared final score. There is no separate victory threshold or individual winner. Record the score and any newly completed Entries in the Steward’s Ledger.')
split('Every placement follows the same sequence:',[
 'To place a Tile, choose an available Tile and a reachable empty hex. Check terrain, adjacency and the Tile’s printed placement rule, pay its cost, then place it and resolve any placement effect.',
 'The Vale has seven terrains: Grasslands, Woodland, Arable Land, Mountains, Heaths, Ruins and Water/River. Use the map legend to identify them. Read each placement requirement exactly: Mine Tunnel goes on Mountains, Farmstead on Arable Land, Salvage Camp on Ruins and Bridge on Water/River. Any Land means an empty non-Water/River hex unless the Tile says otherwise.',
 'Shelter in the Ruins: when placing or upgrading a Tile whose entire footprint lies on Ruins, ignore all Stone in that cost. If any part of its footprint lies outside Ruins, pay the full printed cost.'
])
split('To resolve Linked Production,',[
 'To resolve Linked Production, check the chosen Resource Tile and every immediately adjacent copy from the same Tile stack. If none has participated this round, the chosen Tile and all immediately adjacent, non-Overstrained copies produce for that one Activate Action. Use each producer’s current side.',
 'Linked Production does not continue outward through a chain. Check only the copies immediately adjacent to the chosen Tile.'
])
replace('Treat every Tile that produced','Mark every Tile that produced as having participated this round. If the chosen Tile or any immediately adjacent matching Tile has already participated, only the chosen Tile produces. A disconnected matching set may still use Linked Production. All participation refreshes at the start of the next round.')
split('Boons help the settlement.',[
 'Boons help the settlement. Read the current Season row and resolve it immediately unless the italic line says to keep the card face up. A kept Boon remains active until it is used or its printed duration ends; then discard it.',
 'Using a face-up Boon costs no Action unless its text says otherwise. Only one Boon may modify a single Action unless a rule says otherwise.',
 'For example, A Little Time adds 1 timer to an active Arrival in Season I, up to 3. Its later Season rows add up to 2 or 3 timers as printed.'
])
split('Resolve a Burden by paying',[
 'Resolve a Burden by paying and performing its printed resolution, usually with an Interact Action, or by using an effect that explicitly resolves it. For example, Old Wounds Reopen asks you to choose Social/Wellbeing Tiles, then pay Herbs or take Strain. Its italic line gives the Action and payment needed to discard it.',
 'If a Burden offers payment or Strain, choose the branch before drawing a Target Card. Draw only for the Strain branch. Resolving the Burden stops future reapplication; Strain it has already placed remains.'
])
split('Arrivals are timed objectives',[
 'Arrivals are timed objectives that unlock Special Tiles. When one is revealed, place it face up with 3 timer tokens, one in each circle. Quiet Quest, for example, requires 4 Goods and 2 Herbs and unlocks Special Tile 1, the Adventurers’ Guild.',
 'To complete an Arrival, meet its Requirements and spend 1 Interact Action before its final timer is removed. Remove its timers, discard the card and set the unlocked Special Tile beside the board. Any player may later place it with a normal Place Action, following its printed placement rule and paying 0 resources.',
 'An activated Inn adds 1 timer to an active Arrival; Dawn Break Inn adds up to 2. An Arrival can never hold more than 3 timers.'
])
split('Supported prevents 1 Strain before',[
 'Supported prevents 1 Strain before you check whether a Tile becomes Overstrained. A Tile either has Supported or does not: passive and single-use Supported never stack.',
 'Common Land gives passive Supported to 1 adjacent Housing Tile; The Pleasance gives it to up to 3. Choose the eligible Housing Tiles when the effect becomes active. Keep those choices until the eligible set changes or the source becomes Overstrained, then choose again when required. You cannot reassign passive Supported in response to incoming Strain.',
 'Supported granted by placement, activation or another effect is single-use unless its rule says it is passive.'
])
split('The Steward’s Ledger tracks achievements',[
 'The Steward’s Ledger tracks achievements across games. Before your group’s first game, choose a name and an unused Roman numeral (I–VI) in the Group Register. Use that numeral for every Entry mark and Game Record.',
 'After scoring, mark NAMED for each newly completed Entry and record the shared final score and those Entries. Each NAMED mark counts once toward Golden milestones; player-count prestige marks are records only. Completed Entries are never spent or reset. An AFTER gate is met only if the group begins the game with the shown number of Named Entries.'
])
replace('Tiles are adjacent when','Tiles are adjacent when their hexes share a side. Touching at a corner is not adjacency. Requirements such as ADJ. HOUSING, ADJ. TRAVEL and ADJ. RIVER use this side-sharing rule. A Tile may be reachable through the settlement network without being adjacent. All side-adjacent, non-Overstrained Tiles form that network.')
# A direct learning route, separate from thematic copy.
anchor=find('Contents');anchor.insert_paragraph_before('Using this rulebook','Heading 2');anchor.insert_paragraph_before('For your first game, follow Setup and Encounter seeding in Sections 2–3, then keep the round sequence and four Actions in Sections 4–5 nearby. During play, use the Quick reference and glossary in Section 17. Golden content is optional; leave it aside for your first game.','Normal')
meta={}
def addfigure(anchor,name,label,title,high):
 p=anchor.insert_paragraph_before();p.add_run().add_picture(str(A/(name+'.png')),width=Mm(160));rid=p._p.xpath('.//a:blip/@r:embed')[0];meta[rid]=(label,title,high)
addfigure(find('2. Setup'),'component-families','COMPONENT GUIDE','Tile and card families',85)
addfigure(find('Each Steward has a Starting Terrain'),'steward-guides','STEWARD GUIDES','Choose a different Steward for each player',40)
# A labelled illustrative calculation; this is not a recorded game or invented lore.
anchor=find('Unlocked but unplaced Special Tiles');anchor.insert_paragraph_before('Scoring example','Heading 2')
anchor.insert_paragraph_before('Suppose your group has 80 Population, 25 Renown from Tiles and effects, two completed Steward Objectives (+30) and one eligible Golden bonus (+5). That gives 140 before penalties. One active Burden (−10), two failed Arrivals (−20) and three Strain tokens (−30) reduce the shared final score to 80.','Normal')
for sec in d.sections:
 for hf in [sec.header,sec.footer]:
  for p in hf.paragraphs:
   for run in p.runs:run.text=run.text.replace('v6.7','v6.8')
d.core_properties.version='6.8';out=R/'output/docx/rulebook-v6_8/The_Quiet_Vale_Main_Rulebook_v6_8.docx';out.parent.mkdir(parents=True,exist_ok=True);d.save(out);(W/'wording-changes.json').write_text(json.dumps(edits,ensure_ascii=False,indent=2))
# Retain the original print system while improving scale, flow and hierarchy.
s=(R/'working/rules-v6_7/build_professional_rulebook.py').read_text().replace('v6_7','v6_8').replace('v6.7','v6.8').replace('The_Quiet_Vale_Main_Rulebook_v6_8_Card_Guide.docx','The_Quiet_Vale_Main_Rulebook_v6_8.docx')
s=s.replace('fontSize=8.8,\n    leading=11.3,','fontSize=9.5,\n    leading=12.3,').replace('fontSize=7.75,\n    leading=9.55,','fontSize=8.2,\n    leading=10.2,').replace('fontSize=7.05,\n    leading=8.4,','fontSize=7.8,\n    leading=9.6,').replace('fontSize=6.95,\n    leading=8.15,','fontSize=7.5,\n    leading=9.0,')
s=s.replace('textColor=ROSE,\n    spaceBefore=8.5,','textColor=FOREST_DARK,\n    spaceBefore=10.0,')
s=s.replace('Paragraph(escape(heading.upper()), H2)','Paragraph(escape(heading), H2)')
s=s.replace("elif \"golden tile\" in left:\n            ratio = 0.5","elif \"golden tile\" in left:\n            ratio = 0.34 if headers[1] == 'final scoring condition' else 0.5")
s=s.replace('card = scaled_rl_image(path, 27 * mm, 50 * mm)','card = scaled_rl_image(path, 31 * mm, 55 * mm)').replace('fontSize=7.0, leading=8.5, spaceAfter=0','fontSize=8.0, leading=10.0, spaceAfter=0').replace('colWidths=[29.5 * mm, panel_w - 29.5 * mm]','colWidths=[33.5 * mm, panel_w - 33.5 * mm]')
s=s.replace('Remove its timers, discard it and match its Reward number to the Special Tile back and set the unlocked Tile beside the board.','Remove its timers and discard the card. Match its Reward number to the Special Tile back, then set that Tile beside the board for later placement.')
# Remove production briefs from the reading edition; genuine photographs remain a separate production task.
start=s.index('                if current_section == "12. Stewards":',s.index('def build_story'));end=s.index('                continue',start);s=s[:start]+s[end:]
start=s.index('                if heading == "Setup overview":',s.index('def build_story'));end=s.index('                last_was_break',start);s=s[:start]+s[end:]
start=s.index('            if current_section == "1. Components" and len(table.columns)',s.index('def build_story'));end=s.index('    filler_specs = [',start);s=s[:start]+s[end:]
# Start longer lessons cleanly; keep the quick reference easy to find.
s=s.replace('elif current_section == "8. Encounter cards":', 'elif current_section in {"2. Setup", "5. Your turn and four Actions", "7. Building the settlement", "8. Encounter cards", "9. Strain and Supported", "11. The Target Deck", "14. Final scoring", "16. Golden content", "17. Quick reference"}:')
s=s.replace('story.append(CondPageBreak(150 * mm))','story.append(PageBreak())')
s=s.replace('story.append(CondPageBreak(54 * mm))','story.append(CondPageBreak(76 * mm))')
# New artwork has real descriptive captions and stays with related explanations.
addition='\nIMAGE_META.update({\n'+''.join(f'    {rid!r}: ({lab!r}, {title!r}, {h} * mm),\n' for rid,(lab,title,h) in meta.items())+'})\n\n'
s=s.replace('def build_story(document:',addition+'def build_story(document:')
s=s.replace('f"<b>{escape(self.label)}</b>  <font color=\'#315846\'>{escape(self.title.upper())}</font>"','f"<b>{escape(self.title)}</b>"')
s=s.replace('"Encounter-led visual redesign proof using locked game content"','"Polished reading edition with current artist tiles and numbered Arrival cards"')
# Make the three allowed signature-padding pages useful references instead of production placeholders.
start=s.index('    filler_specs = [',s.index('def build_story'));end=s.index('    story.append(NextPageTemplate("BackCover"))',start)
s=s[:start]+'''    if filler_count:
        from polish_reference import reference_pages
        story.extend(reference_pages(filler_count, globals()))

'''+s[end:]
(W/'build_professional_rulebook.py').write_text(s)
# The fidelity audit checks the edited source plus unchanged mechanical requirements.
s=(R/'working/rules-v6_7/audit_professional_rulebook.py').read_text().replace('v6_7','v6_8')
for label in ['Complete component family photograph','Final overhead setup photograph','Six Steward role products and role-linked markers',"Final settlement and Steward's Ledger close-up"]:s=s.replace('        '+repr(label).replace("'",'"')+',\n','') if "'" not in label else s.replace('        "'+label+'",\n','')
(W/'audit_professional_rulebook.py').write_text(s)
print('Wrote polished v6.8 source and print builder;',len(edits),'wording edits',meta)
