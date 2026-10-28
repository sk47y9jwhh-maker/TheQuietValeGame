from pathlib import Path
from docx import Document
R=Path(__file__).resolve().parents[2];W=R/'working/rules-v6_8';p=R/'output/docx/rulebook-v6_8/The_Quiet_Vale_Main_Rulebook_v6_8.docx';d=Document(p)
for par in list(d.paragraphs):
 if par.text in ['Setup continued','Encounter types and resolution']:par._p.getparent().remove(par._p)
d.save(p)
p=W/'build_professional_rulebook.py';s=p.read_text()
for old,new in {'INTERACT WITH AN ENCOUNTER':'Interact with an Encounter','SEASON TRANSITION':'Season transition','KEY COSTS AND LIMITS':'Key costs and limits','WHERE TO LOOK':'Where to look','THE VALE ACROSS THREE SEASONS':'The Vale across three Seasons','GOLDEN BOONS - OPTIONAL ADVANCED EFFECTS':'Golden Boons — optional advanced effects'}.items():s=s.replace('Paragraph("'+old+'", H2)','Paragraph("'+new+'", H2)')
s=s.replace('else:\n                    story.append(CondPageBreak(76 * mm))','elif current_section == "15. Steward\'s Ledger":\n                    story.append(CondPageBreak(40 * mm))\n                else:\n                    story.append(CondPageBreak(76 * mm))')
s=s.replace('                story.append(Paragraph(escape(heading), H2))','                if heading in {"Build the Encounter material", "Terrain, placement and adjacency"}:\n                    story.append(CondPageBreak(75 * mm))\n                story.append(Paragraph(escape(heading), H2))')
s=s.replace('                for rel_id in rel_ids:\n                    path = rel_paths.get(rel_id)','                for rel_id in rel_ids:\n                    if rel_id == "rId48":\n                        continue  # Golden faces are shown at readable size beside their scoring conditions.\n                    path = rel_paths.get(rel_id)')
# Resolve the Golden scoring reference as one full illustrated page.
s=s.replace('def styled_table(table: DocxTable, rel_paths:', 'def styled_table(table: DocxTable, rel_paths:')
needle='    if len(table.columns) == 1:\n        return one_cell_callout(table.cell(0, 0).text)'
s=s.replace(needle,'''    if len(table.columns) == 2 and normalise_spaces(table.cell(0, 1).text) == "Final scoring condition":
        from polish_reference import golden_scoring_gallery
        return golden_scoring_gallery(table, globals())
'''+needle)
# Start the Golden table on a fresh page under its own clear title.
needle='            rendered = styled_table(table, rel_paths, crossrefs)'
if needle not in s:
 print('Rendered table location differs; inspect')
else:s=s.replace(needle,'''            if len(table.columns) == 2 and normalise_spaces(table.cell(0, 1).text) == "Final scoring condition":
                story.extend([PageBreak(), SectionHeader("Golden Tile scoring", include_in_toc=False), Spacer(1, 4 * mm), Paragraph("Each placed, non-Overstrained Golden Tile earns +5 Renown at final scoring only if its condition is met.", BODY)])
'''+needle)
# The figure caption is intentionally replaced by the full gallery heading.
s=s.replace("'Five optional Ledger rewards'","'Golden Tile scoring'")
needle='''            font_size=6.7,
        ),
    ]'''
s=s.replace(needle,'''            font_size=6.7,
        ),
        Spacer(1, 4 * mm),
        *special_number_lookup(),
    ]''',1)
s=s.replace('def quick_reference_supplement(','''def special_number_lookup():
    from polish_reference import special_number_table
    return special_number_table(globals())


def quick_reference_supplement(''')
p.write_text(s)
p=W/'polish_reference.py';s=p.read_text().replace("Paragraph('Tile categories',g['H2'])","Paragraph('Category and Golden symbols',g['H2'])").replace("('Gold','Golden')","('Gold','Golden Tile')")
s+='''

def golden_scoring_gallery(table,g):
    data=[]
    ids=['charter','hearth','river_gate','cairn','garden']
    for row,ident in zip(table.rows[1:],ids):
        art=g['scaled_rl_image'](R/f'output/artist-tiles-complete/tiles-web/golden_tile_the_golden_{ident}.png',36*mm,32*mm)
        content=[Paragraph(escape(row.cells[0].text),g['H2']),Paragraph(escape(row.cells[1].text),g['BODY'])]
        data.append([art,content])
    output=Table(data,colWidths=[43*mm,g['BODY_W']-43*mm],rowHeights=[36*mm]*5,hAlign='LEFT')
    output.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('LINEBELOW',(0,0),(-1,-1),.3,g['LINE'])]))
    return output


def special_number_table(g):
    with (R/'design/tile-set/special-back-manifest.csv').open() as f:rows=list(csv.DictReader(f))
    pairs=[]
    for i in range(13):
        a,b=rows[i],rows[i+13]
        pairs.append((a['special_tile_number'],a['paired_tile_name'],b['special_tile_number'],b['paired_tile_name']))
    return [Paragraph('Special Tile numbers',g['H2']),Paragraph('Match the Arrival Reward number to its Tile back. Stables uses both 17 and 18.',g['BODY_SMALL']),g['compact_reference_table'](['No.','Special Tile','No.','Special Tile'],pairs,[9*mm,g['BODY_W']/2-9*mm,9*mm,g['BODY_W']/2-9*mm],font_size=7.0)]
'''
p.write_text(s)
p=W/'audit_professional_rulebook.py';s=p.read_text().replace('Passive Supported cannot be reassigned in response to incoming Strain','You cannot reassign passive Supported in response to incoming Strain').replace('approved_visual_labels = [f"{label} {title}" for label, title, _ in builder.IMAGE_META.values()]','approved_visual_labels = [title for label, title, _ in builder.IMAGE_META.values()]');p.write_text(s)
