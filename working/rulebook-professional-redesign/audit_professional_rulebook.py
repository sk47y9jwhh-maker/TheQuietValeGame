from __future__ import annotations

import importlib.util
import json
import re
import sys
import unicodedata
from pathlib import Path

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[2]
BUILDER_PATH = ROOT / "working/rulebook-professional-redesign/build_professional_rulebook.py"
REPORT_PATH = ROOT / "tmp/pdfs/rulebook-v6_6/fidelity-audit.json"


def load_builder():
    spec = importlib.util.spec_from_file_location("quiet_vale_rulebook_builder", BUILDER_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError("Could not load rulebook builder")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def normalise(value: str) -> str:
    text = unicodedata.normalize("NFKC", value)
    text = (
        text.replace("’", "'")
        .replace("‘", "'")
        .replace("“", '"')
        .replace("”", '"')
        .replace("\u00a0", " ")
    )
    return " ".join(text.split())


def canonical(value: str) -> str:
    return " ".join(re.findall(r"[a-z0-9]+(?:'[a-z0-9]+)?|[+>-]", normalise(value).casefold()))


def fragment_present(fragment: str, combined_canonical: str) -> bool:
    expected = canonical(fragment)
    if expected in combined_canonical:
        return True
    tokens = expected.split()
    if len(tokens) < 12:
        return False
    windows = [" ".join(tokens[index : index + 8]) for index in range(0, len(tokens) - 7, 6)]
    if not windows:
        return False
    matched = sum(window in combined_canonical for window in windows)
    return matched / len(windows) >= 0.85


def embedded_font_records(reader: PdfReader) -> list[dict[str, object]]:
    standard_14 = {
        "/Courier", "/Courier-Bold", "/Courier-Oblique", "/Courier-BoldOblique",
        "/Helvetica", "/Helvetica-Bold", "/Helvetica-Oblique", "/Helvetica-BoldOblique",
        "/Times-Roman", "/Times-Bold", "/Times-Italic", "/Times-BoldItalic",
        "/Symbol", "/ZapfDingbats",
    }
    records: dict[str, dict[str, object]] = {}
    for page in reader.pages:
        resources = page.get("/Resources") or {}
        fonts = resources.get("/Font") or {}
        for resource_name, reference in fonts.items():
            font = reference.get_object()
            base_font = str(font.get("/BaseFont", resource_name))
            descriptor = font.get("/FontDescriptor")
            if descriptor is None and font.get("/DescendantFonts"):
                descendant = font["/DescendantFonts"][0].get_object()
                descriptor = descendant.get("/FontDescriptor")
            embedded = False
            if descriptor is not None:
                descriptor = descriptor.get_object()
                embedded = any(key in descriptor for key in ("/FontFile", "/FontFile2", "/FontFile3"))
            records[base_font] = {
                "base_font": base_font,
                "subtype": str(font.get("/Subtype", "")),
                "embedded": embedded,
                "standard_14": base_font in standard_14,
            }
    return sorted(records.values(), key=lambda item: str(item["base_font"]))


def source_fragments(builder, encounter_page: int) -> list[str]:
    document = Document(builder.SOURCE)
    crossrefs = {"encounters": encounter_page}
    fragments: list[str] = []
    cover_finished = False
    for kind, block in builder.iter_blocks(document):
        if not cover_finished:
            if kind == "P" and builder.paragraph_has_page_break(block):
                cover_finished = True
            continue
        if kind == "P":
            value = normalise(builder.transform_text(block.text, crossrefs))
            if value.startswith("After Rounds 4 and 8, resolve these steps in order"):
                # The professional layout expresses this source paragraph as
                # the equivalent six-row Season transition checklist.
                continue
            if value:
                fragments.append(value)
            continue
        if (
            len(block.columns) == 2
            and builder.normalise_spaces(block.cell(0, 0).text) == "Section"
            and builder.normalise_spaces(block.cell(0, 1).text) == "Page"
        ):
            continue
        for row in block.rows:
            for cell in row.cells:
                value = normalise(builder.transform_text(cell.text, crossrefs))
                if value:
                    fragments.append(value)
    return fragments


def main() -> None:
    builder = load_builder()
    reader = PdfReader(str(builder.OUTPUT))
    page_texts = [normalise(page.extract_text() or "") for page in reader.pages]
    combined = normalise(" ".join(page_texts))
    combined_canonical = canonical(combined)
    section_pages = builder.scan_section_pages(builder.OUTPUT, builder.source_section_titles(Document(builder.SOURCE)))
    encounter_page = section_pages.get("8. Encounter cards", -1)

    fragments = source_fragments(builder, encounter_page)
    missing_fragments = [fragment for fragment in fragments if len(fragment) >= 12 and not fragment_present(fragment, combined_canonical)]

    required = [
        "Shuffle all 30 Target Cards and place the deck face down.",
        "63 double-sided Core Tiles across 21 Tile types",
        "26 one-sided Special Tile components across 25 designs",
        "90 Encounter Cards: 30 Boons, 30 Burdens, 25 Arrivals and 5 Golden Boons",
        "90 colour-coded resource pieces: 15 each of Wood, Stone, Metal, Food, Herbs and Goods",
        "30 black markers and 30 bronze markers",
        "Arrival timers may use either colour",
        "1 Main Rulebook, 6 Steward Aids and 6 matching Steward Tokens",
        "1 Steward's Ledger booklet for tracking achievements, scores and Golden milestones across multiple games",
        "5 optional one-sided Golden Tiles",
        "take 4 Boons, 4 Burdens and 4 Arrivals",
        "12-card-per-player pool",
        "Choose the player order",
        "PENDING EFFECT",
        "Pay the resource cost printed for the upgraded side",
        "INTERACT WITH AN ENCOUNTER",
        "A 0-Action benefit saves the Action only",
        "An unlocked Special Tile costs 0 resources to place",
        "Build the eligible pool from all placed Tiles below 3 Strain",
        "If no such Tile exists, place no Strain",
        "MULTIPLE FAILED ARRIVALS",
        "PARTIAL RESOLUTION",
        "Fewer active Burdens than the player count",
        "SEASON TRANSITION",
        "THE VALE ACROSS THREE SEASONS",
        "KEY COSTS AND LIMITS",
        "Available Tile",
        "Pending effect",
        "Open card lets the players choose any remaining Tile",
        "There is no separate victory threshold",
        "Each player chooses one Steward, takes its Steward Aid and Steward Token",
        "Linked Production does not continue outward through a chain",
        "Passive Supported",
        "Single-use Supported",
        "Passive Supported cannot be reassigned in response to incoming Strain",
        "Before your group's first game, choose a name and an unused Roman numeral (I-VI) in the Group Register",
        "After scoring, mark NAMED for each newly completed Entry",
        "An AFTER gate is met only if the group begins the game with the shown number of Named Entries",
        "A group may attempt only one Vow per game",
        "Completed Entries are never spent or reset",
        "Record the shared final score and any newly completed Entries in the Steward's Ledger",
        "Shelter in the Ruins",
        "When placing or upgrading a Tile whose entire footprint lies on Ruins, ignore all Stone in that cost",
        "TWO TIMING DECISIONS",
        "2 standard cards per player remain in the deck after Season I",
        "This is your clearest control over the Season row",
        "Hold or place it lower",
        "It still begins with 3 timers when revealed",
        "Subtract 10 for each Arrival that expired during play",
        "Subtract 10 for every Strain token on the map",
        "A group reaches Golden milestones at 5, 12, 18, 25 and 32 Named Entries",
        "One roof became three. A path reached the workshop, another reached the river, and neighbours carried tools between them. Nobody planned the whole shape alone. The settlement grew wherever one piece of work made the next possible.",
        "After the rain, one wall bowed and a path washed thin. Before either failed, neighbours arrived with braces, stone, and dry timber. Nothing grand was done that day. The settlement held because ordinary repairs were made in time.",
        "The result is a board that feels increasingly lived-in. By the end of play, the settlement is not just a score state; it is a record of the group’s choices.",
        "Credits, edition and rules support",
        "Complete component family photograph",
        "Final overhead setup photograph",
        "Six Steward role products and role-linked markers",
        "Final settlement and Steward's Ledger close-up",
        "PRODUCTION PLACEHOLDER - ICON REFERENCE",
    ]
    forbidden = [
        "24 Target Cards",
        "24-card deck",
        "take 5 Boons, 5 Burdens and 5 Arrivals",
        "Place upgraded sides",
        "No active Burdens",
        "The group chooses the order in which players act",
        "USE THE COMPANION RULES REFERENCE v6.4",
        "Encounter card guide on page 10",
        "online prototype",
        "Round Timer",
        "River/Water",
        "Steward Player Aid",
        "one persistent spread",
        "Use one spread for one group",
        "Final scorepad",
        "approved scorepad",
        "final total and any completed",
        "Golden unlocks",
        "FIRST-ROUND WALKTHROUGH",
        "FIRST GAME - TEACH IN LAYERS",
        "TEACHING THE FIRST GAME",
        "TEACHING ADVICE",
        "first teaching game",
        "WHY SEEDING MATTERS",
        "SEED WITH INTENT",
        "Seed the Encounter Deck",
        "earlier to gain more rounds to complete it",
        "Unfulfilled Promise",
        "Subtract 5 for each",
        "MORE TERMS",
        "the resolver voluntarily chose zero targets",
        "COSTS, LIMITS AND MOVEMENT",
        "ENCOUNTER-LED DESIGN PROOF",
        "BOON / MANY HANDS, LIGHT WORK",
        "BOON / FROM THE BRINK",
        "BOON / SHARED HANDS",
        "□",
        "�",
    ]

    required_results = {item: canonical(item) in combined_canonical for item in required}
    approved_visual_labels = [f"{label} {title}" for label, title, _ in builder.IMAGE_META.values()]
    approved_visual_results = {
        item: canonical(item) in combined_canonical
        for item in approved_visual_labels
    }
    forbidden_results = {
        item: (canonical(item) in combined_canonical if canonical(item) else item in combined)
        for item in forbidden
    }
    trim_width_mm = float(reader.pages[0].mediabox.width) / 72 * 25.4
    trim_height_mm = float(reader.pages[0].mediabox.height) / 72 * 25.4
    sparse_pages = [index for index, text in enumerate(page_texts, 1) if 1 < index < len(page_texts) and len(text) < 45]
    font_records = embedded_font_records(reader)
    fonts_embedded = bool(font_records) and all(
        bool(record["embedded"]) or bool(record["standard_14"])
        for record in font_records
    )

    report = {
        "pdf": str(builder.OUTPUT),
        "page_count": len(reader.pages),
        "signature_multiple_of_four": len(reader.pages) % 4 == 0,
        "trim_mm": [round(trim_width_mm, 3), round(trim_height_mm, 3)],
        "section_pages": section_pages,
        "encounter_cross_reference_page": encounter_page,
        "source_fragment_count": len(fragments),
        "missing_source_fragments": missing_fragments,
        "required_checks": required_results,
        "approved_visual_checks": approved_visual_results,
        "forbidden_checks": forbidden_results,
        "sparse_interior_pages": sparse_pages,
        "fonts": font_records,
        "all_custom_fonts_embedded": fonts_embedded,
        "pass": (
            len(reader.pages) % 4 == 0
            and abs(trim_width_mm - 195) < 0.05
            and abs(trim_height_mm - 255) < 0.05
            and not missing_fragments
            and all(required_results.values())
            and all(approved_visual_results.values())
            and not any(forbidden_results.values())
            and not sparse_pages
            and fonts_embedded
        ),
    }
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))
    if not report["pass"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
