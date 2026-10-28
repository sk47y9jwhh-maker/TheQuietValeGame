from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
DOCX = ROOT / "output/docx/rulebook-v6_4/The_Quiet_Vale_Main_Rulebook_v6_4_Card_Guide.docx"
TARGET_DIR = ROOT / "output/target-card-visual"
VISUAL = ROOT / "working/rulebook-v6_4-card-visuals/target-card-in-action.png"

INK = "25251F"
GREEN = "365745"
BRASS = "9B7A38"
PALE = "F5F0E2"
PALE_ALT = "ECE6D5"
CALLOUT = "F4E7BB"
GRID = "857A62"
TABLE_DXA = 9920


def set_font(run, *, size: float, bold: bool = False, italic: bool = False, color: str = INK) -> None:
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, *, top: int = 90, start: int = 130, bottom: int = 90, end: int = 130) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_table_borders(table, *, color: str = GRID, size: int = 7) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def format_p(paragraph, *, align=WD_ALIGN_PARAGRAPH.LEFT, after=0, line=1.0) -> None:
    paragraph.style = "Normal"
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.keep_together = True


def set_picture_alt(inline_shape, description: str) -> None:
    inline_shape._inline.docPr.set("descr", description)
    inline_shape._inline.docPr.set("title", description)


def add_target_card_cell(cell, *, image_path: Path, label: str, body: str, alt: str, fill: str) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, top=100, start=155, bottom=100, end=155)
    set_cell_shading(cell, fill)

    image_p = cell.paragraphs[0]
    format_p(image_p, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    shape = image_p.add_run().add_picture(str(image_path), width=Inches(1.48))
    set_picture_alt(shape, alt)

    label_p = cell.add_paragraph()
    format_p(label_p, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    set_font(label_p.add_run(label), size=9.6, bold=True, color=GREEN)

    body_p = cell.add_paragraph()
    format_p(body_p, after=0, line=1.03)
    set_font(body_p.add_run(body), size=8.9, color=INK)


def add_heading(document: Document, text: str):
    paragraph = document.add_paragraph(text, style="Heading 2")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(document: Document, text: str, *, after: float = 4):
    paragraph = document.add_paragraph(text, style="Normal")
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.keep_together = True
    return paragraph


def add_example_table(document: Document):
    table = document.add_table(rows=1, cols=2)
    set_table_geometry(table, [2700, 7220])
    set_table_borders(table, color=GRID, size=7)
    headers = [("READ THIS LINE", GREEN), ("WHAT REMAINS", GREEN)]
    for cell, (text, fill) in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, fill)
        set_cell_margins(cell, top=75, start=120, bottom=75, end=120)
        p = cell.paragraphs[0]
        format_p(p, after=0)
        set_font(p.add_run(text), size=8.7, bold=True, color="FFFFFF")

    rows = [
        ("Eligible pool", "Improved Path; Fine Cabin; Path; Alms House"),
        ("Tile Class → Core", "Improved Path, Fine Cabin and Path remain. Alms House is removed."),
        ("Tile Side → Upgraded", "Improved Path and Fine Cabin remain. Path is removed."),
        ("Adjacency → 3+ Neighbours", "Improved Path and Fine Cabin both remain."),
        ("Current Strain → Already Strained", "Both still match, so neither is removed."),
        ("Tie Direction → SE", "Compare the remaining Tiles with each other in the printed SE direction. Improved Path is most southeast, so it is selected."),
    ]
    for index, (left, right) in enumerate(rows):
        cells = table.add_row().cells
        fill = PALE if index % 2 == 0 else PALE_ALT
        for cell in cells:
            set_cell_shading(cell, fill)
            set_cell_margins(cell, top=65, start=120, bottom=65, end=120)
        p = cells[0].paragraphs[0]
        format_p(p, after=0, line=1.0)
        set_font(p.add_run(left), size=8.5, bold=True, color=GREEN)
        p = cells[1].paragraphs[0]
        format_p(p, after=0, line=1.0)
        set_font(p.add_run(right), size=8.5, color=INK)
    return table


def add_callout(document: Document):
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_DXA])
    set_table_borders(table, color=BRASS, size=9)
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    set_cell_margins(cell, top=90, start=145, bottom=90, end=145)
    p = cell.paragraphs[0]
    format_p(p, after=0, line=1.04)
    set_font(p.add_run("FINISH  "), size=9.0, bold=True, color=GREEN)
    set_font(p.add_run("Return the card to the bottom of the Target Card deck. Then resolve Supported and other prevention. Prevented Strain is not redirected."), size=9.0, color=INK)
    return table


def add_visual(document: Document):
    if not VISUAL.exists():
        raise FileNotFoundError(VISUAL)
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_together = True
    shape = paragraph.add_run().add_picture(str(VISUAL), width=Inches(6.82))
    set_picture_alt(
        shape,
        "Target Card 05 example: the eligible pool contains Improved Path, Fine Cabin, Path and Alms House; Core removes Alms House, Upgraded removes Path, and the SE direction selects Improved Path over Fine Cabin.",
    )
    return paragraph


def remove_between(start, end) -> None:
    node = start._p.getnext()
    while node is not end._p:
        next_node = node.getnext()
        node.getparent().remove(node)
        node = next_node


def move_after(anchor, element):
    anchor.addnext(element)
    return element


def main() -> None:
    document = Document(DOCX)
    reachability = next(
        p for p in document.paragraphs
        if p.text.strip().startswith("A Tile or empty hex is reachable")
    )
    reachability.text = (
        "A Tile or empty hex is reachable when the acting Steward can reach it through the shared connected settlement network, or when a card, Tile or Steward rule makes it reachable. All Stewards may use the shared network. Your Steward Token marks your current position in that network. After your first placement, whenever you Place, Upgrade or Activate, move the Token onto the Tile you just placed, upgraded or activated; your next reachability check starts from that Tile."
    )
    heading = next(p for p in document.paragraphs if p.text.strip() == "11. The Target Deck")
    next_heading = next(p for p in document.paragraphs if p.text.strip() == "12. Stewards")
    next_heading.paragraph_format.page_break_before = True
    remove_between(heading, next_heading)

    structured = TARGET_DIR / "target-card-05-vale-blue-rose-front.png"
    open_card = TARGET_DIR / "target-card-25-open-vale-blue-rose-front.png"
    for path in (structured, open_card):
        if not path.exists():
            raise FileNotFoundError(path)

    elements = []
    elements.append(add_body(document, "When an effect adds Strain, first build the eligible pool. A Target Card then chooses from that pool; it never adds a Tile that the effect did not make eligible.", after=5))

    cards = document.add_table(rows=1, cols=2)
    set_table_geometry(cards, [TABLE_DXA // 2, TABLE_DXA // 2])
    set_table_borders(cards, color=GRID, size=7)
    add_target_card_cell(
        cards.cell(0, 0),
        image_path=structured,
        label="STRUCTURED • CARD 05",
        body="Read top to bottom: Core → Upgraded → 3+ Neighbours → Already Strained. If more than one Tile survives, compare those Tiles in the printed SE direction and select the one most southeast.",
        alt="Locked Target Card 05, structured Vale Blue and Rose front",
        fill=PALE,
    )
    add_target_card_cell(
        cards.cell(0, 1),
        image_path=open_card,
        label="OPEN • CARD 25",
        body="Every preference says Any. Keep the eligible pool, then choose any remaining eligible Tile. No arrow is used.",
        alt="Locked Target Card 25, open Vale Blue and Rose front",
        fill=PALE_ALT,
    )
    elements.append(cards)
    elements.append(add_visual(document))
    elements.append(add_callout(document))

    anchor = heading._p
    for item in elements:
        element = item._p if hasattr(item, "_p") else item._tbl
        anchor = move_after(anchor, element)

    document.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
