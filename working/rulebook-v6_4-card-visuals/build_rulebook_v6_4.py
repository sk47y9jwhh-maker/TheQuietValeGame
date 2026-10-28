from __future__ import annotations

import hashlib
import re
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / "working/rulebook-v6_4-card-visuals"
SOURCE = WORK / "source/The_Quiet_Vale_Main_Rulebook_v6_3_OneDrive_Table_Fix.docx"
OUTPUT = ROOT / "output/docx/rulebook-v6_5/The_Quiet_Vale_Main_Rulebook_v6_5_Card_Guide.docx"
CARD_DIR = WORK / "current-card-examples"
EXPECTED_SOURCE_SHA256 = "92ab91319694c0d452eac0752d5b1dab909157c209fad7cc9aace3f0091614ec"

INK = "25251F"
GREEN = "365745"
BRASS = "9B7A38"
PALE = "F5F0E2"
PALE_ALT = "ECE6D5"
CALLOUT = "F4E7BB"
GRID = "857A62"
TABLE_DXA = 9920


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def paragraph_by_text(document: Document, text: str):
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.strip() == text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph {text!r}; found {len(matches)}")
    return matches[0]


def replace_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_repeatable_run_font(run, *, size: float, bold: bool = False, italic: bool = False, color: str = INK) -> None:
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, *, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_table_borders(table, *, color: str = GRID, size: int = 8) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths: list[int]) -> None:
    if len(widths) != len(table.columns):
        raise ValueError("Width count must match table column count")
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_picture_alt(inline_shape, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description)


def format_cell_paragraph(paragraph, *, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0, line=1.0) -> None:
    paragraph.style = "Normal"
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.keep_together = True


def add_card_to_cell(cell, *, image_path: Path, label: str, map_line: str, explanation: str, alt: str) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, top=85, start=135, bottom=95, end=135)

    image_p = cell.paragraphs[0]
    format_cell_paragraph(image_p, after=2)
    picture_run = image_p.add_run()
    shape = picture_run.add_picture(str(image_path), width=Inches(1.22))
    set_picture_alt(shape, alt)

    label_p = cell.add_paragraph()
    format_cell_paragraph(label_p, after=1)
    label_run = label_p.add_run(label)
    set_repeatable_run_font(label_run, size=10.0, bold=True, color=GREEN if label != "GOLDEN BOON" else BRASS)

    map_p = cell.add_paragraph()
    format_cell_paragraph(map_p, after=2, line=1.0)
    map_run = map_p.add_run(map_line)
    set_repeatable_run_font(map_run, size=8.7, bold=True, color=INK)

    explanation_p = cell.add_paragraph()
    format_cell_paragraph(explanation_p, align=WD_ALIGN_PARAGRAPH.LEFT, after=0, line=1.04)
    explanation_run = explanation_p.add_run(explanation)
    set_repeatable_run_font(explanation_run, size=8.7, color=INK)


def build_card_grid(document: Document):
    table = document.add_table(rows=2, cols=2)
    set_table_geometry(table, [TABLE_DXA // 2, TABLE_DXA // 2])
    set_table_borders(table, color=GRID, size=7)
    for index, row in enumerate(table.rows):
        keep_row_together(row)
        for cell in row.cells:
            set_cell_shading(cell, PALE if index == 0 else PALE_ALT)

    cards = [
        (
            table.cell(0, 0),
            CARD_DIR / "boon-a-little-time.png",
            "BOON",
            "Season numeral → effect → italic duration",
            "Use the row matching the current Season. The lower italic line says whether the card is immediate, remains face up, has limited uses or expires.",
            "Boon card example: A Little Time",
        ),
        (
            table.cell(0, 1),
            CARD_DIR / "burden-old-wounds-reopen.png",
            "BURDEN",
            "Season numeral → effect → italic resolution",
            "Apply the current row when revealed and leave the card active. Its lower italic line gives the Action and cost used to resolve and discard it.",
            "Burden card example: Old Wounds Reopen",
        ),
        (
            table.cell(1, 0),
            CARD_DIR / "arrival-quiet-quest.png",
            "ARRIVAL",
            "Requirements → Reward → three timer circles",
            "Complete its Requirements with an Interact Action. The Reward unlocks a Special Tile for later placement. Put one timer token in each circle when revealed.",
            "Arrival card example: Quiet Quest",
        ),
        (
            table.cell(1, 1),
            CARD_DIR / "golden-boon-the-golden-scroll.png",
            "GOLDEN BOON",
            "Single Effect → italic instruction",
            "Golden Boons have no Season rows or timer circles. They resolve free, do not count toward the standard reveal total, and never enter hands or seeding.",
            "Golden Boon card example: The Golden Scroll",
        ),
    ]
    for cell, image_path, label, map_line, explanation, alt in cards:
        if not image_path.exists():
            raise FileNotFoundError(image_path)
        add_card_to_cell(
            cell,
            image_path=image_path,
            label=label,
            map_line=map_line,
            explanation=explanation,
            alt=alt,
        )
    return table


def build_callout(document: Document, label: str, body: str):
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_DXA])
    set_table_borders(table, color=BRASS, size=9)
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    set_cell_margins(cell, top=105, start=150, bottom=105, end=150)
    p = cell.paragraphs[0]
    format_cell_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.08)
    label_run = p.add_run(f"{label}  ")
    set_repeatable_run_font(label_run, size=9.4, bold=True, color=GREEN)
    body_run = p.add_run(body)
    set_repeatable_run_font(body_run, size=9.4, color=INK)
    return table


def update_contents(document: Document) -> None:
    table = document.tables[3]
    entries = [
        ("Components and setup", "3–5"),
        ("Seasons, seeding and rounds", "6–7"),
        ("Actions and building", "8–13"),
        ("Encounter cards and card guide", "14–16"),
        ("Strain and Target Cards", "16–19"),
        ("Stewards", "20"),
        ("Final scoring and Ledger", "22"),
        ("Golden content", "23"),
        ("Quick reference and glossary", "24–25"),
    ]
    if len(table.rows) != 10:
        raise RuntimeError("Unexpected contents-table row count")
    for row, (section, page) in zip(table.rows[1:], entries):
        replace_paragraph_text(row.cells[0].paragraphs[0], section)
        replace_paragraph_text(row.cells[1].paragraphs[0], page)


def update_footer_version(document: Document) -> None:
    replacements = 0
    for section in document.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            for paragraph in footer.paragraphs:
                for run in paragraph.runs:
                    if "v6.3" in run.text:
                        run.text = run.text.replace("v6.3", "v6.5")
                        replacements += 1
    if replacements == 0:
        raise RuntimeError("Footer version text was not found")


def request_field_refresh(document: Document) -> None:
    settings = document.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def update_five_player_rules(document: Document) -> None:
    replace_paragraph_text(
        paragraph_by_text(document, "1–4 players  •  45–120 minutes  •  Cooperative"),
        "1–5 players  •  45–150 minutes  •  Cooperative",
    )

    steward_choice = paragraph_by_text(
        document,
        "Each player chooses a different Steward, takes its Player Aid and Steward Token, and reads its Starting Terrain, once-per-Season Power and Steward Objective aloud. There are no Starting Benefits.",
    )
    replace_paragraph_text(
        steward_choice,
        "Each player chooses a different Steward, takes its Steward Aid and Steward Token, and reads its Starting Terrain, once-per-Season Power and Steward Objective aloud. There are no Starting Benefits. In a 1-player game, choose 1 Steward and use the 1-player values wherever shown. In a 5-player game, choose 5 of the 6 Stewards; the unchosen Steward is not used.",
    )

    replace_paragraph_text(
        paragraph_by_text(
            document,
            "2.5  For each player, take 5 Boons, 5 Burdens and 5 Arrivals to form a balanced Encounter pool.",
        ),
        "2.5 Separate the standard Encounter Cards by type and shuffle each type. For each player, take 4 Boons, 4 Burdens and 4 Arrivals; combine all selected cards into one 12-card-per-player pool. In a 5-player game, this is 20 of each type, 60 cards total.",
    )
    replace_paragraph_text(
        paragraph_by_text(document, "2.6  Shuffle that pool and deal 9 hidden Encounter Cards to each player."),
        "2.6 Shuffle the combined selected pool. Deal 9 hidden Encounter Cards to each player.",
    )
    replace_paragraph_text(
        paragraph_by_text(document, "2.7  Place 3 standard Encounter Cards per player into the starting Encounter Deck."),
        "2.7 Without adding any other cards, place the remaining 3 selected standard Encounter Cards per player face down as the starting Encounter Deck. At 5 players, deal 45 cards and place the remaining 15 in the deck.",
    )
    replace_paragraph_text(
        paragraph_by_text(
            document,
            "2.8  Keep the remaining standard cards face down as the reserve. Players may not inspect them unless an effect instructs them to do so.",
        ),
        "2.8 Return every unselected standard Encounter Card to the box as the reserve. Keep it separate from the starting deck; players may not inspect it unless an effect instructs them to do so.",
    )

    warehouse_table = next(
        (
            table
            for table in document.tables
            if table.rows
            and len(table.columns) == 2
            and table.cell(0, 0).text.strip() == "Players"
            and table.cell(0, 1).text.strip() == "Start with"
        ),
        None,
    )
    if warehouse_table is None:
        raise RuntimeError("Starting Warehouse table was not found")
    player_rows = {row.cells[0].text.strip(): row for row in warehouse_table.rows[1:]}
    if "5" not in player_rows:
        new_row = deepcopy(warehouse_table.rows[-1]._tr)
        warehouse_table._tbl.append(new_row)
        player_rows = {row.cells[0].text.strip(): row for row in warehouse_table.rows[1:]}
        player_rows["4"].cells[0].text = "5"
        player_rows["4"].cells[1].text = "0 resources"

    five_row = warehouse_table.rows[-1]
    replace_paragraph_text(five_row.cells[0].paragraphs[0], "5")
    replace_paragraph_text(five_row.cells[1].paragraphs[0], "0 resources")


def set_section_rule_numbering(document: Document) -> None:
    """Make rule lists restart within each numbered main section.

    The source document uses Word's single continuous ``List Number`` style,
    which makes the visible rule numbers run across unrelated sections. Use
    literal section.rule labels instead so the numbering is stable in both
    Word and the rendered PDF: 2.1, 2.2 … then 3.1, 3.2, and so on.
    """
    current_section = None
    counters: dict[int, int] = {}

    for paragraph in document.paragraphs:
        heading_match = re.match(r"^(\d+)\.", paragraph.text.strip()) if paragraph.style.name == "Heading 1" else None
        if heading_match:
            current_section = int(heading_match.group(1))
            counters.setdefault(current_section, 0)

        if paragraph.style.name != "List Number":
            continue
        if current_section is None:
            raise RuntimeError(f"Found a numbered rule before a numbered section: {paragraph.text!r}")

        counters[current_section] += 1
        label = f"{current_section}.{counters[current_section]}"
        original_text = paragraph.text.strip()
        replace_paragraph_text(paragraph, f"{label}  {original_text}")
        paragraph.style = "Normal"
        paragraph.paragraph_format.left_indent = Inches(0.34)
        paragraph.paragraph_format.first_line_indent = Inches(-0.34)
        paragraph.paragraph_format.space_after = Pt(2.5)


def main() -> None:
    actual_sha = sha256(SOURCE)
    if actual_sha != EXPECTED_SOURCE_SHA256:
        raise RuntimeError(f"Source master changed: {actual_sha}")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)

    set_section_rule_numbering(document)
    update_five_player_rules(document)
    update_contents(document)
    update_footer_version(document)
    request_field_refresh(document)
    document.core_properties.title = "The Quiet Vale — Main Rulebook v6.5"
    document.core_properties.subject = "Complete rules with Encounter Card visual guide"

    encounter_heading = paragraph_by_text(document, "8. Encounter Cards")
    replace_paragraph_text(encounter_heading, "8. Encounter cards")
    encounter_heading.paragraph_format.page_break_before = True

    boons_heading = paragraph_by_text(document, "Boons")

    intro = boons_heading.insert_paragraph_before(
        "Every Encounter is read from top to bottom. The title and italic flavour text identify its story; the rules appear in the printed sections below. Specific card text overrides a general rule only for the situation it names.",
        style="Normal",
    )
    intro.paragraph_format.keep_together = True

    reading_heading = boons_heading.insert_paragraph_before("Encounter types at a glance", style="Heading 2")
    reading_heading.paragraph_format.keep_with_next = True

    card_grid = build_card_grid(document)
    boons_heading._p.addprevious(card_grid._tbl)

    current_season = build_callout(
        document,
        "CURRENT SEASON",
        "Use I in Foundations (Rounds 1–4), II in Growth (Rounds 5–8), and III in Maturity (Rounds 9–12). The online prototype highlights the current row; on a physical card, follow the matching Roman numeral.",
    )
    boons_heading._p.addprevious(current_season._tbl)

    italic_note = boons_heading.insert_paragraph_before(
        "Flavour text is narrative only. The separate italic line inside a Boon or Burden Season row is rules text; on a Golden Boon, the italic line at the base is rules text.",
        style="Normal",
    )
    italic_note.paragraph_format.space_before = Pt(4)
    italic_note.paragraph_format.keep_together = True

    resolving_heading = boons_heading.insert_paragraph_before(
        "Encounter types and resolution", style="Heading 2"
    )
    resolving_heading.paragraph_format.page_break_before = True
    resolving_heading.paragraph_format.keep_with_next = True

    replace_paragraph_text(boons_heading, "Boons — temporary help")

    boon_rule = paragraph_by_text(document, "Resolve the effect for the current Season. Most Boons resolve immediately. A deferred Boon remains face up until used or until its printed duration ends. Only one Boon may modify a single Action unless a rule says otherwise.")
    replace_paragraph_text(
        boon_rule,
        "Boons help the settlement and are resolved from the row matching the current Season. Resolve the effect immediately unless the italic line says to keep the card face up. A kept Boon remains active until it is used or its printed duration ends, then it is discarded. For example, A Little Time adds 1 timer to an active Arrival in Season I, up to 3; its Season II and III rows add up to 2 or 3 timers as printed. Only one Boon may modify a single Action unless a rule says otherwise.",
    )

    burdens_heading = paragraph_by_text(document, "Burdens")
    replace_paragraph_text(burdens_heading, "Burdens — ongoing pressure")
    burden_apply = paragraph_by_text(
        document,
        "Apply the current-Season effect, then place the card face up as an active Burden. At the start of Seasons II and III, every unresolved Burden reapplies using the new Season effect.",
    )
    replace_paragraph_text(
        burden_apply,
        "Burdens create ongoing pressure. Apply the current-Season effect when the card is revealed, then place it face up as an active Burden. At the start of Seasons II and III, every unresolved Burden reapplies using the new Season effect.",
    )
    burden_resolution = paragraph_by_text(document, "Resolve a Burden by paying and performing its printed resolution, normally with an Interact Action, or by using an effect that explicitly resolves it. Resolution prevents later reapplication but does not remove Strain already placed.")
    replace_paragraph_text(
        burden_resolution,
        "Resolve a Burden by paying and performing its printed resolution, normally with an Interact Action, or by using an effect that explicitly resolves it. For example, Old Wounds Reopen lets you choose Social/Wellbeing Tiles and either pay Herbs or take Strain; its italic line gives the Action and payment needed to discard it. If a Burden offers payment or Strain, choose the branch before drawing a Target Card; draw only if the Strain branch is chosen. Resolution prevents later reapplication but does not remove Strain already placed.",
    )

    arrivals_heading = paragraph_by_text(document, "Arrivals")
    replace_paragraph_text(arrivals_heading, "Arrivals — timed objectives")
    arrival_rule = paragraph_by_text(document, "Place a revealed Arrival face up with 3 timer tokens. To complete it, fulfil its Requirement and spend 1 Interact Action. Remove its timers and place its unlocked Special Tile on the card. Any player may later place that Tile using a normal Place Action and its printed placement rule.")
    replace_paragraph_text(
        arrival_rule,
        "Arrivals are timed objectives that unlock Special Tiles. Place a revealed Arrival face up with 3 timer tokens, one in each circle. For example, Quiet Quest requires 4 Goods and 2 Herbs and rewards the Adventurers’ Guild. To complete it, fulfil its Requirements and spend 1 Interact Action before the final timer is removed. Remove its timers, discard the Arrival and set its unlocked Special Tile beside the board; completing an Arrival unlocks its Tile but does not place it immediately. Timers, expiry and Inn activations are shown in the Arrival examples below.",
    )

    section9 = paragraph_by_text(document, "9. Strain and Supported")
    golden_heading = section9.insert_paragraph_before("Golden Boons — optional advanced effects", style="Heading 2")
    golden_heading.paragraph_format.keep_with_next = True
    golden_body = section9.insert_paragraph_before(
        "Golden Boons are optional advanced cards unlocked through the Steward’s Ledger. They never enter player hands or seasonal seeding. When a Golden Boon is revealed, resolve it for free, do not count it toward the standard reveal total, and continue revealing. The Golden Scroll, for example, lets each player discard any standard Encounter Cards from their hand, then draw the same number at random from the box; Golden Boons cannot be drawn, and if the box runs short, draw as many as possible. Resolve the card, then discard it.",
        style="Normal",
    )
    golden_body.paragraph_format.keep_together = True

    golden_last = paragraph_by_text(document, "Use each Golden component’s printed rule. Every Golden Tile has a conditional +5 Renown scoring hook.")
    replace_paragraph_text(
        golden_last,
        "Use each Golden component’s printed rule. Every Golden Tile has a conditional +5 Renown scoring hook. The Golden Boon format is shown in the Encounter card guide on page 14: read its single Effect block, then follow the italic instruction at the base.",
    )

    companion_reference = paragraph_by_text(
        document,
        "FOR DEFINITIONS, TIMING AND EDGE CASES, USE THE COMPANION RULES REFERENCE v6.3.",
    )
    replace_paragraph_text(
        companion_reference,
        "FOR DEFINITIONS, TIMING AND EDGE CASES, USE THE COMPANION RULES REFERENCE v6.5.",
    )

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
