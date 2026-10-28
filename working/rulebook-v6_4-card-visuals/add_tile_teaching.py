from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
DOCX = ROOT / "output/docx/rulebook-v6_4/The_Quiet_Vale_Main_Rulebook_v6_4_Card_Guide.docx"
VISUAL_DIR = ROOT / "working/rulebook-v6_4-card-visuals"

INK = "25251F"
GREEN = "365745"


def set_font(run, *, size: float, bold: bool = False, color: str = INK) -> None:
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_picture_alt(shape, description: str) -> None:
    shape._inline.docPr.set("descr", description)
    shape._inline.docPr.set("title", description)


def replace_cover_image() -> None:
    """Replace the source cover artwork while preserving the DOCX package."""
    cover = VISUAL_DIR / "cover-hero.png"
    temporary = DOCX.with_name(DOCX.stem + ".cover.tmp.docx")
    with ZipFile(DOCX, "r") as source, ZipFile(temporary, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = cover.read_bytes() if item.filename == "word/media/image1.png" else source.read(item.filename)
            target.writestr(item, data)
    temporary.replace(DOCX)


def add_body(document: Document, text: str):
    p = document.add_paragraph(text, style="Normal")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_together = True
    return p


def add_heading(document: Document, text: str):
    p = document.add_paragraph(text, style="Heading 2")
    p.paragraph_format.keep_with_next = True
    return p


def add_visual(document: Document, filename: str, alt: str, width: float = 6.82):
    path = VISUAL_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    p = document.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_together = True
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    set_picture_alt(shape, alt)
    return p


def insert_after(anchor, element):
    node = element._p if hasattr(element, "_p") else element._tbl
    anchor_node = anchor._p if hasattr(anchor, "_p") else anchor._tbl if hasattr(anchor, "_tbl") else anchor
    anchor_node.addnext(node)
    return node


def insert_before(anchor, element):
    node = element._p if hasattr(element, "_p") else element._tbl
    anchor_node = anchor._p if hasattr(anchor, "_p") else anchor._tbl if hasattr(anchor, "_tbl") else anchor
    anchor_node.addprevious(node)
    return node


def paragraph_start(document: Document, prefix: str):
    return next(p for p in document.paragraphs if p.text.strip().startswith(prefix))


def already_added(document: Document, marker: str) -> bool:
    if any(marker in p.text for p in document.paragraphs):
        return True
    return any(
        marker in (shape._inline.docPr.get("descr") or "")
        for shape in document.inline_shapes
    )


def remove_paragraph(paragraph) -> None:
    node = paragraph._element
    node.getparent().remove(node)


def remove_drawings_between(document: Document, start_prefix: str, stop_prefix: str) -> None:
    active = False
    for paragraph in list(document.paragraphs):
        text = paragraph.text.strip()
        if text.startswith(start_prefix):
            active = True
            continue
        if active and text.startswith(stop_prefix):
            break
        if active and paragraph._p.xpath(".//w:drawing"):
            remove_paragraph(paragraph)


def table_containing(document: Document, text: str):
    for table in document.tables:
        contents = "\n".join(cell.text for row in table.rows for cell in row.cells)
        if text in contents:
            return table
    raise RuntimeError(f"Table containing {text!r} was not found")


def replace_table_text(document: Document, old: str, new: str) -> None:
    table = table_containing(document, old)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                if old in paragraph.text:
                    paragraph.text = paragraph.text.replace(old, new)


def main():
    document = Document(DOCX)

    # Strengthen the action language beside the visual examples.
    upgrade = paragraph_start(document, "Spend 1 Action and pay the listed cost to flip a placed basic Core Tile")
    upgrade.text = (
        "Spend 1 Action and pay the listed cost to flip a placed basic Core Tile to its upgraded side. It remains the same placed Tile; only its name, cost, effect and printed scoring values change. Upgraded, Special, Golden and Overstrained Tiles cannot be upgraded."
    )

    # Replace the legacy setup, season and action diagrams with the unified
    # teaching set. These are kept in their own sections so each visual has a
    # single job and can be read alongside the relevant rule text.
    remove_drawings_between(document, "2. Setup", "3. Seasons and Encounter seeding")
    setup_overview = paragraph_start(document, "Setup overview")
    if not already_added(document, "Standard Setup visual"):
        setup_visual = add_visual(
            document,
            "standard-setup.png",
            "Standard Setup visual showing the starting map, the three shuffled decks, the resource supply, and each player's starting Steward Token and Encounter hand.",
            width=6.82,
        )
        insert_after(setup_overview._p, setup_visual)

    remove_drawings_between(document, "3. Seasons and Encounter seeding", "5. Your turn and four Actions")
    round_heading = paragraph_start(document, "4. Round sequence")
    round_heading.paragraph_format.page_break_before = True
    if not already_added(document, "Seasons and Rounds visual"):
        seasons_visual = add_visual(
            document,
            "seasons-rounds.png",
            "Seasons and Rounds visual showing Foundations, Growth and Maturity across twelve rounds, with Season changes after Rounds 4 and 8.",
            width=6.82,
        )
        insert_after(round_heading._p, seasons_visual)

    seed_complete = paragraph_start(document, "Complete this before Rounds 1, 5 and 9")
    if not already_added(document, "Encounter Seeding visual"):
        seeding_visual = add_visual(
            document,
            "encounter-seeding.png",
            "Encounter Seeding visual showing each player secretly choosing three hidden cards, then placing one on top, one into the middle without counting, and one on the bottom of the Encounter Deck.",
            width=6.82,
        )
        insert_after(seed_complete._p, seeding_visual)

    remove_drawings_between(document, "5. Your turn and four Actions", "6. Production and the Warehouse")
    action_intro = paragraph_start(document, "On your turn, spend up to 4 Actions")
    if not already_added(document, "Four Actions visual"):
        actions_visual = add_visual(
            document,
            "four-actions.png",
            "Four Actions visual showing Place, Upgrade, Activate and Interact as the four available Action types on a turn.",
            width=6.82,
        )
        insert_after(action_intro._p, actions_visual)

    activate = paragraph_start(document, "Spend 1 Action to use a reachable, non-Overstrained Tile with an Activated Effect")
    if not already_added(document, "Steward Token movement"):
        movement_intro = add_body(
            document,
            "Your Steward Token follows the Tile you most recently changed. After the first placement, Place, Upgrade or Activate a Tile, move the Token onto that Tile; the next reachability check starts there.",
        )
        movement_visual = add_visual(
            document,
            "steward-token-movement.png",
            "Three real board-state examples showing a Steward Token after placing Cabin, upgrading to Fine Cabin, and activating Workshops.",
            width=6.82,
        )
        insert_after(activate._p, movement_intro)
        insert_after(movement_intro._p, movement_visual)

    # Keep Section 6's examples in Section 6, then place the river example in
    # Section 7. The source master previously merged these into one image.
    section7 = paragraph_start(document, "7. Building the settlement")
    linked_text = paragraph_start(document, "Further activations involving that group during the same round")
    if not already_added(document, "Linked Production visual"):
        linked_visual = add_visual(
            document,
            "linked-production.png",
            "Linked Production using three connected upgraded Mine Shaft Tiles: the first activation produces 15 resources, while a later activation produces only from the chosen Tile.",
            width=6.82,
        )
        linked_visual.runs[0]._r.get_or_add_rPr()
        linked_visual._p.getparent().remove(linked_visual._p)
        linked_text._p.addnext(linked_visual._p)
        replace_table_text(
            document,
            "Three adjacent upgraded matching Resource Tiles each produce 5.",
            "Three adjacent upgraded Mine Shaft Tiles each produce 5.",
        )

    remove_drawings_between(document, "7. Building the settlement", "8. Encounter cards")
    adjacency = paragraph_start(document, "Adjacency and connection")
    paragraphs = list(document.paragraphs)
    adjacency_index = next(i for i, paragraph in enumerate(paragraphs) if paragraph._p is adjacency._p)
    adjacency_body = paragraphs[adjacency_index + 1]
    adjacency_body.text = (
        "Two Tiles are adjacent when their hexes share a side. Placement markers such as ADJ. HOUSING, ADJ. TRAVEL and ADJ. RIVER refer to this side-sharing relationship. Adjacency is not the same as reachability: a Tile may be reachable through the connected settlement network without being adjacent. All side-adjacent, non-Overstrained Tiles form the connected settlement network."
    )
    if not already_added(document, "Terrain, placement and adjacency visual"):
        terrain_heading = add_heading(document, "Terrain, placement and adjacency")
        terrain_body = add_body(
            document,
            "Every Tile placement passes the same checks: choose an available Tile, choose an empty reachable hex, check its terrain, adjacency and printed placement rule, pay the listed cost, then place the Tile and resolve any placement effect. The Vale contains Grasslands, Woodland, Arable Land, Mountains, Heaths, Ruins and River/Water terrain. Use the map terrain legend below to identify each region before applying a placement rule. Read the Tile’s printed requirement exactly: Mine Shaft goes on Mountains, Farmstead on Arable Land, Salvage Camp on Ruins and Bridge on Water/River. Ruins are special terrain: they provide a Stone discount. Any Land means an empty non-River hex unless the Tile’s own rule says otherwise.",
        )
        terrain_visual = add_visual(
            document,
            "terrain-placement.png",
            "Terrain, placement and adjacency visual showing real Mine Shaft, Farmstead, Salvage Camp and Bridge terrain examples; side-sharing adjacency examples with Path, Workshops, Cabin and Cottage; and the five placement checks.",
            width=6.82,
        )
        insert_before(adjacency._p, terrain_heading)
        insert_before(adjacency._p, terrain_body)
        insert_before(adjacency._p, terrain_visual)
        terrain_legend = add_visual(
            document,
            "terrain-map-legend.png",
            "Map terrain legend using clipped examples from the real map: Mountains, Woodland, Arable Land, Heaths, Ruins, River/Water and Grasslands, each labelled for matching map artwork to the placement rules.",
            width=6.82,
        )
        insert_before(adjacency._p, terrain_legend)

    rivers = paragraph_start(document, "Rivers are barriers.")
    if not already_added(document, "Rivers and Bridges visual"):
        river_visual = add_visual(
            document,
            "river-connection.png",
            "Separate river examples showing that opposite banks do not connect without a Bridge, and that a Bridge on Water/River reconnects the settlement network.",
            width=6.82,
        )
        insert_after(rivers._p, river_visual)

    section7.paragraph_format.page_break_before = True
    if not already_added(document, "Workshops: basic and upgraded"):
        workshops_heading = add_heading(document, "Workshops: basic and upgraded")
        workshops_body = add_body(
            document,
            "Workshops: Once per round, when upgrading an adjacent Core Tile, reduce its cost by 1 resource. Its upgraded side, The Makers Conclave, changes the range: once per round, when upgrading a reachable Core Tile, reduce its cost by up to 2 resources.",
        )
        workshops_visual = add_visual(
            document,
            "workshops-core.png",
            "Separate Workshops and The Makers Conclave examples showing adjacency on the basic side and reachable range on the upgraded side.",
            width=6.82,
        )
        insert_before(section7._p, workshops_heading)
        insert_before(section7._p, workshops_body)
        insert_before(section7._p, workshops_visual)

        merchant_heading = add_heading(document, "Market Stalls: converting Goods")
        merchant_body = add_body(
            document,
            "Market Stalls: Once per round, when paying a cost, you may spend 1 Goods as 1 resource of any type. Its upgraded side, The Seldes, lets you spend 1 Goods as up to 2 resources of one type. The printed Goods requirement is never replaced by this conversion.",
        )
        merchant_visual = add_visual(
            document,
            "merchant-core.png",
            "Separate Market Stalls and The Seldes examples showing Goods conversions for the Fine Cabin Tile cost and the Quiet Quest Arrival cost.",
            width=6.82,
        )
        merchant_heading.paragraph_format.page_break_before = True
        insert_before(section7._p, merchant_heading)
        insert_before(section7._p, merchant_body)
        insert_before(section7._p, merchant_visual)

    # Replace the legacy combined Overstrain example with a dedicated visual.
    remove_drawings_between(document, "10. Overstrain chains", "A Tile already at 3")
    if not already_added(document, "Overstrain Chain visual"):
        chain_table = table_containing(document, "CHAIN EXAMPLE")
        overstrain_visual = add_visual(
            document,
            "overstrain-chain.png",
            "Overstrain Chain using black Strain discs: Cottage reaches three Strain, a fresh Target Card selects the next eligible Tile, and the chain continues only when a Tile crosses to three.",
            width=6.82,
        )
        insert_before(chain_table, overstrain_visual)
    replace_table_text(
        document,
        "Tile A and adjacent Tile B each have 2 Strain. An effect adds 1 to A. After the effect finishes, A spreads to the eligible adjacent Tile selected by a Target Card. If that Tile is B, B reaches 3 and later makes its own spread.",
        "Cottage and adjacent Path each have 2 Strain. An effect adds 1 to Cottage. After the effect finishes, Cottage spreads to the eligible adjacent Tile selected by a Target Card. If that Tile is Path, Path reaches 3 and later makes its own spread.",
    )

    # Make the Arrival section explicitly teach Inn timing.
    arrival = paragraph_start(document, "Arrivals are timed objectives that unlock Special Tiles")
    arrival.text = (
        "Arrivals are timed objectives that unlock Special Tiles. Place a revealed Arrival face up with 3 timer tokens, one in each circle. For example, Quiet Quest requires 4 Goods and 2 Herbs and rewards the Adventurers’ Guild. To complete it, fulfil its Requirements and spend 1 Interact Action before the final timer is removed. Remove its timers, discard the Arrival and set its unlocked Special Tile beside the board; completing an Arrival unlocks its Tile but does not place it immediately. Timers, expiry and Inn activations are shown in the Arrival examples below. An Inn adds 1 timer to an active Arrival when activated; an upgraded Dawn Break Inn adds up to 2, and no Arrival can hold more than 3 timers."
    )
    remove_drawings_between(document, "8. Encounter cards", "9. Strain and Supported")
    if not already_added(document, "Arrival Timing visual"):
        timing_visual = add_visual(
            document,
            "arrival-timing.png",
            "Arrival Timing visual showing three timers revealed, one timer removed at the end of each round, completion while a timer remains, and failure when the final timer is removed.",
            width=6.82,
        )
        insert_after(arrival._p, timing_visual)
    if not already_added(document, "Inn and Arrival timers"):
        golden_heading = paragraph_start(document, "Golden Boons — optional advanced effects")
        inn_visual = add_visual(
            document,
            "inn-arrival-timers.png",
            "Inn and Dawn Break Inn examples showing timer loss at the end of a round and an Inn activation restoring an active Quiet Quest Arrival from 1 to 2 timers.",
            width=6.82,
        )
        insert_before(golden_heading._p, inn_visual)

    # Make Core Tile support concrete in the Supported section.
    supported = paragraph_start(document, "Supported prevents Strain before checking")
    supported.text = (
        "Supported prevents Strain before checking whether a Tile becomes Overstrained. A Tile either has Supported or does not; multiple instances do not stack. Some Core Tiles provide passive Supported: Common Land protects 1 adjacent Housing Tile, while its upgraded side, The Pleasance, protects up to 3. The source must not be Overstrained."
    )
    prevention = paragraph_start(document, "The Golden Garden and any other prevention")
    if not already_added(document, "A Core Tile giving Supported"):
        support_visual = add_visual(
            document,
            "supported-from-core.png",
            "Common Land and The Pleasance giving passive Supported to adjacent Housing Tiles, with the upgraded side supporting up to three.",
            width=6.82,
        )
        support_visual.paragraph_format.page_break_before = True
        insert_before(prevention._p, support_visual)

    remove_drawings_between(document, "13. End of a Season", "14. Final scoring")
    section13 = paragraph_start(document, "13. End of a Season")
    section13.paragraph_format.page_break_before = True
    support_rounds = paragraph_start(document, "At the end of Seasons I and II")
    if not already_added(document, "Neighbourly Support visual"):
        neighbour_visual = add_visual(
            document,
            "neighbourly-support.png",
            "Neighbourly Support visual showing a Housing cluster receiving one Supported token for every three eligible non-Overstrained Housing Tiles, while an Overstrained Tile is ignored.",
            width=6.82,
        )
        insert_after(support_rounds._p, neighbour_visual)

    document.save(DOCX)
    replace_cover_image()
    print(DOCX)


if __name__ == "__main__":
    main()
