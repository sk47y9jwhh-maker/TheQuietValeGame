from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (
    ROOT
    / "output/docx/rulebook-v6_5/The_Quiet_Vale_Main_Rulebook_v6_5_Card_Guide.docx"
)
OUTPUT = (
    ROOT
    / "output/docx/rulebook-v6_6/The_Quiet_Vale_Main_Rulebook_v6_6_Card_Guide.docx"
)
VISUAL_ROOT = ROOT / "working/rulebook-v6_4-card-visuals"
VISUAL_REPLACEMENTS = {
    "word/media/image17.png": VISUAL_ROOT / "standard-setup.png",
    "word/media/image18.png": VISUAL_ROOT / "seasons-rounds.png",
    "word/media/image19.png": VISUAL_ROOT / "encounter-seeding.png",
    "word/media/image20.png": VISUAL_ROOT / "four-actions.png",
    "word/media/image21.png": VISUAL_ROOT / "steward-token-movement.png",
    "word/media/image22.png": VISUAL_ROOT / "linked-production.png",
    "word/media/image23.png": VISUAL_ROOT / "terrain-placement.png",
    "word/media/image24.png": VISUAL_ROOT / "terrain-map-legend.png",
    "word/media/image25.png": VISUAL_ROOT / "river-connection.png",
    "word/media/image26.png": VISUAL_ROOT / "workshops-core.png",
    "word/media/image27.png": VISUAL_ROOT / "merchant-core.png",
    "word/media/image28.png": VISUAL_ROOT / "overstrain-chain.png",
    "word/media/image29.png": VISUAL_ROOT / "arrival-timing.png",
    "word/media/image30.png": VISUAL_ROOT / "inn-arrival-timers.png",
    "word/media/image31.png": VISUAL_ROOT / "supported-from-core.png",
    "word/media/image32.png": VISUAL_ROOT / "neighbourly-support.png",
}


def compact(value: str) -> str:
    return " ".join(value.split())


def replace_run_text(paragraph: Paragraph, old: str, new: str) -> None:
    replacements = 0
    for run in paragraph.runs:
        if old not in run.text:
            continue
        run.text = run.text.replace(old, new)
        replacements += 1
    if replacements != 1:
        raise RuntimeError(
            f"Expected one run replacement for {old!r}; found {replacements}"
        )


def set_paragraph_text(paragraph: Paragraph, value: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def set_labeled_paragraph(paragraph: Paragraph, label: str, body: str) -> None:
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = f"{label}  "
        first.bold = True
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        first = paragraph.add_run(f"{label}  ")
        first.bold = True
    paragraph.add_run(body)


def paragraph_by_text(document: Document, value: str) -> Paragraph:
    wanted = compact(value)
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if compact(paragraph.text) == wanted
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph {wanted!r}; found {len(matches)}")
    return matches[0]


def replace_paragraph(document: Document, old: str, new: str) -> None:
    set_paragraph_text(paragraph_by_text(document, old), new)


def cell_paragraph_by_text(document: Document, value: str) -> Paragraph:
    wanted = compact(value)
    matches: list[Paragraph] = []
    seen: set[object] = set()
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                marker = cell._tc
                if marker in seen:
                    continue
                seen.add(marker)
                if compact(cell.text) == wanted and len(cell.paragraphs) == 1:
                    matches.append(cell.paragraphs[0])
    if len(matches) != 1:
        raise RuntimeError(f"Expected one table cell {wanted!r}; found {len(matches)}")
    return matches[0]


def replace_cell(document: Document, old: str, new: str) -> None:
    set_paragraph_text(cell_paragraph_by_text(document, old), new)


def table_containing(document: Document, value: str) -> Table:
    matches = [
        table
        for table in document.tables
        if any(value in cell.text for row in table.rows for cell in row.cells)
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one table containing {value!r}; found {len(matches)}")
    return matches[0]


def remove_row(table: Table, row_index: int) -> None:
    table._tbl.remove(table.rows[row_index]._tr)


def remove_equal_width_column(table: Table, column_index: int) -> None:
    grid = table._tbl.tblGrid
    widths = [
        int(column.get(qn("w:w")) or "0")
        for column in list(grid)
    ]
    if len(widths) != len(table.columns):
        raise RuntimeError("Penalty table grid does not match its visible columns")
    total_width = sum(widths)

    for row in table.rows:
        row._tr.remove(row.cells[column_index]._tc)
    grid.remove(list(grid)[column_index])

    remaining_columns = list(grid)
    new_width = total_width // len(remaining_columns)
    for index, column in enumerate(remaining_columns):
        width = new_width if index < len(remaining_columns) - 1 else total_width - new_width * index
        column.set(qn("w:w"), str(width))
        for row in table.rows:
            tc_width = row.cells[index]._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_width is not None:
                tc_width.set(qn("w:w"), str(width))
                tc_width.set(qn("w:type"), "dxa")


def remove_heading_section(document: Document, heading: str) -> None:
    body = document.element.body
    removing = False
    found = False
    for child in list(body.iterchildren()):
        if child.tag == qn("w:sectPr"):
            continue
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            if compact(paragraph.text) == compact(heading):
                removing = True
                found = True
            elif removing and paragraph.style.name == "Heading 1":
                break
        if removing:
            body.remove(child)
    if not found:
        raise RuntimeError(f"Section heading was not found: {heading}")


def remove_preceding_page_break(
    document: Document,
    *,
    paragraph_text: str | None = None,
    relationship_id: str | None = None,
) -> None:
    if (paragraph_text is None) == (relationship_id is None):
        raise ValueError("Specify exactly one page-break target")
    targets: list[Paragraph] = []
    for paragraph in document.paragraphs:
        text_matches = (
            paragraph_text is not None
            and compact(paragraph.text) == compact(paragraph_text)
        )
        rel_ids = {
            blip.get(qn("r:embed"))
            for blip in paragraph._p.xpath(".//a:blip")
        }
        rel_matches = relationship_id is not None and relationship_id in rel_ids
        if text_matches or rel_matches:
            targets.append(paragraph)
    if len(targets) != 1:
        label = paragraph_text or relationship_id
        raise RuntimeError(f"Expected one page-break target {label!r}; found {len(targets)}")
    target = targets[0]._p
    previous = target.getprevious()
    if previous is None or not previous.xpath('.//w:br[@w:type="page"]'):
        raise RuntimeError("The target is not preceded by a hard page break")
    previous.getparent().remove(previous)


def remove_paragraph_with_relationship(document: Document, relationship_id: str) -> None:
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if relationship_id
        in {
            blip.get(qn("r:embed"))
            for blip in paragraph._p.xpath(".//a:blip")
        }
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph with relationship {relationship_id!r}; found {len(matches)}"
        )
    paragraph = matches[0]._p
    paragraph.getparent().remove(paragraph)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    existing = properties.find(qn("w:shd"))
    if existing is not None:
        properties.remove(existing)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_table_widths(table: Table, widths: list[int]) -> None:
    if len(widths) != len(table.columns):
        raise ValueError("Table width count does not match its columns")
    table.autofit = False
    grid_columns = list(table._tbl.tblGrid)
    for column_index, width in enumerate(widths):
        grid_columns[column_index].set(qn("w:w"), str(width))
        for row in table.rows:
            cell_width = row.cells[column_index]._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                row.cells[column_index]._tc.get_or_add_tcPr().append(cell_width)
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def write_decision_cell(cell, text: str, *, header: bool, shaded: bool) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Aptos")
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Aptos")
        run.font.size = Pt(9 if header else 8.5)
        run.bold = header
        if header:
            run.font.color.rgb = None
            color = run._element.get_or_add_rPr().get_or_add_color()
            color.set(qn("w:val"), "FFFFFF")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "243E32" if header else ("F8F4EA" if shaded else "FFFFFF"))


def add_seeding_decision_table(document: Document, after_table: Table) -> None:
    rows = [
        (
            "Boon",
            "Accelerate a plan or give a kept effect more time before Season end.",
            "Use a stronger later-Season row, or wait until its targets exist.",
        ),
        (
            "Burden",
            "Reveal its weaker current-Season row while resolution is cheaper and more time remains before reapplication.",
            "Gain preparation time, but risk a stronger row, a higher resolution cost and less response time.",
        ),
        (
            "Arrival",
            "Unlock its Special Tile with more game left to place and use it. It still begins with 3 timers when revealed.",
            "Wait until its Requirement is easier to meet, but leave less time to benefit from the Tile.",
        ),
    ]
    table = document.add_table(rows=1 + len(rows), cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [1450, 4271, 4272])
    headers = ("Card", "Bring it forward", "Hold or place it lower")
    for column_index, value in enumerate(headers):
        write_decision_cell(table.cell(0, column_index), value, header=True, shaded=False)
    for row_index, values in enumerate(rows, 1):
        for column_index, value in enumerate(values):
            write_decision_cell(
                table.cell(row_index, column_index),
                value,
                header=False,
                shaded=row_index % 2 == 0,
            )
    after_table._tbl.addnext(table._tbl)


def update_footer_versions(document: Document) -> None:
    replacements = 0
    for section in document.sections:
        for footer in (
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            for paragraph in footer.paragraphs:
                for run in paragraph.runs:
                    if "v6.5" in run.text:
                        run.text = run.text.replace("v6.5", "v6.6")
                        replacements += 1
    if replacements == 0:
        raise RuntimeError("No v6.5 footer version was found")


def replace_docx_members(docx_path: Path, replacements: dict[str, Path]) -> None:
    for replacement in replacements.values():
        if not replacement.is_file():
            raise FileNotFoundError(replacement)
    temporary = docx_path.with_suffix(".tmp.docx")
    found: set[str] = set()
    with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename in replacements:
                data = replacements[info.filename].read_bytes()
                found.add(info.filename)
            target.writestr(info, data)
    missing = set(replacements) - found
    if missing:
        temporary.unlink(missing_ok=True)
        raise RuntimeError(f"DOCX members were not found: {sorted(missing)}")
    temporary.replace(docx_path)


def apply_editorial_sweep(document: Document) -> None:
    paragraph_replacements = {
        "The banners of war have long since fallen silent. Roads once marched by armies now welcome traders, settlers and wandering folk searching for somewhere safe to begin again. Across three Seasons, you will gather resources, expand one shared settlement, welcome Arrivals and overcome the Burdens left behind by a fractured world.":
            "War’s banners have long since fallen silent. Old military roads now carry traders, settlers and wandering folk looking for somewhere safe to begin again. Across three Seasons, you will gather resources, grow one shared settlement, welcome Arrivals and meet the Burdens of a fractured world.",
        "The Quiet Vale is not a game about conquest. It is a game about stewardship, resilience and building a place that may endure.":
            "The Quiet Vale is a game of stewardship: build carefully, support your neighbours and leave the Vale stronger than you found it.",
        "Work together to build Population and earn Renown across 12 rounds. Complete Steward Objectives, fulfil Arrivals, resolve Burdens and protect the settlement from Strain. At the end of Round 12, total your shared score and record the story of the Vale you created.":
            "Together, build Population and earn Renown over 12 rounds. Complete Steward Objectives, fulfil Arrivals, resolve Burdens and protect the settlement from Strain. After Round 12, calculate one shared final score. There is no separate victory threshold; the score records what your group achieved. Record it, together with any newly completed Entries, in the Steward’s Ledger.",
        "2.1 Lay out the Game Map. Place the Stewards Board, Warehouse and Round Timer beside it.":
            "2.1  Lay out the Game Map. Place the Stewards Board and Warehouse beside it.",
        "2.2 Organise the basic Core Tiles with their basic sides visible. Place upgraded sides, Special Tiles and tokens within reach.":
            "2.2  Organise the Core Tile supply with every basic side face up. Each upgraded side is the reverse of the same physical Tile; do not make a separate upgraded supply. Place Special Tiles and tokens within reach.",
        "2.3 Place the Round marker on Round 1.":
            "2.3  Place the Round marker on Round 1 of the round tracker.",
        "Each player chooses a different Steward, takes its Steward Aid and Steward Token, and reads its Starting Terrain, once-per-Season Power and Steward Objective aloud. There are no Starting Benefits. In a 1-player game, choose 1 Steward and use the 1-player values wherever shown. In a 5-player game, choose 5 of the 6 Stewards; the unchosen Steward is not used.":
            "Each player chooses a different Steward and takes the matching Steward Aid and Steward Token. Read each Steward’s Starting Terrain, once-per-Season Power and Objective aloud. Stewards have no Starting Benefits. For 1 player, choose 1 Steward and use the 1-player values wherever shown. For 5 players, choose any 5 and return the sixth Steward to the box.",
        "2.5 Separate the standard Encounter Cards by type and shuffle each type. For each player, take 4 Boons, 4 Burdens and 4 Arrivals; combine all selected cards into one 12-card-per-player pool. In a 5-player game, this is 20 of each type, 60 cards total.":
            "2.5  Separate the standard Encounter Cards by type, then shuffle each type. For each player, take 4 Boons, 4 Burdens and 4 Arrivals. Combine them into one 12-card-per-player pool. For 5 players, take 20 of each type: 60 cards total.",
        "2.6 Shuffle the combined selected pool. Deal 9 hidden Encounter Cards to each player.":
            "2.6  Shuffle the selected pool, then deal 9 hidden Encounter Cards to each player.",
        "2.7 Without adding any other cards, place the remaining 3 selected standard Encounter Cards per player face down as the starting Encounter Deck. At 5 players, deal 45 cards and place the remaining 15 in the deck.":
            "2.7  Place the remaining 3 selected cards per player face down as the starting Encounter Deck. Add no other cards. For 5 players, deal 45 cards and place the remaining 15 in the deck.",
        "2.8 Return every unselected standard Encounter Card to the box as the reserve. Keep it separate from the starting deck; players may not inspect it unless an effect instructs them to do so.":
            "2.8  Return every unselected standard Encounter Card to the box as the reserve. Keep the reserve separate from the starting deck and do not inspect it unless an effect instructs you to do so.",
        "Each player places their Steward Token on an empty hex matching their Steward’s Starting Terrain. This costs no Actions or resources and does not place a Tile.":
            "Each player places their Steward Token on an empty hex matching that Steward’s Starting Terrain. This places no Tile and costs no Actions or resources.",
        "The game has three Seasons of four rounds each: Foundations (Rounds 1–4), Growth (5–8) and Maturity (9–12).":
            "A game lasts 12 rounds across three four-round Seasons: Foundations (Rounds 1–4), Growth (Rounds 5–8) and Maturity (Rounds 9–12). Before each Season, every player commits 3 cards from their hand and chooses where each enters the Encounter Deck.",
        "Seed at the start of every Season":
            "Choose the Season, then the position",
        "3.1 Each player secretly chooses 3 cards from their hand.":
            "3.1  Choose 3 cards from your hidden hand. In Seasons I and II, keep every unchosen card in your hand for a later Season.",
        "3.2 Seed one card on top of the Encounter Deck.":
            "3.2  Seed one chosen card on top of the Encounter Deck.",
        "3.3 Insert one card into the middle without counting the deck.":
            "3.3  Insert one chosen card into the middle without counting the deck.",
        "3.4 Seed one card on the bottom.":
            "3.4  Seed the remaining chosen card on the bottom.",
        "Complete this before Rounds 1, 5 and 9. Once a card is seeded, nobody may count, inspect, reorder or reveal cards in the deck unless an effect allows it.":
            "Seed before Rounds 1, 5 and 9. Cards already in the deck stay there. In normal play, 2 standard cards per player remain in the deck after Season I, 1 per player remains after Season II, and every remaining standard card is revealed in Season III. A lower card may therefore cross into the next Season and use that Season’s row. Nobody may count, inspect, reorder or reveal the deck unless an effect allows it.",
        "On your turn, spend up to 4 Actions. You may repeat an Action if every use is legal. Unused Actions are lost when your turn ends.":
            "On your turn, take up to 4 Actions. You may repeat an Action whenever each use is legal. Unused Actions are lost when your turn ends.",
        "No Tile may be placed on River/Water unless its own rule permits or requires it. Bridges follow their printed River placement rule.":
            "Do not place a Tile on Water/River terrain unless its own rule permits or requires it. A Bridge follows its printed Water/River placement rule.",
        "Spend 1 Action and pay the listed cost to flip a placed basic Core Tile to its upgraded side. It remains the same placed Tile; only its name, cost, effect and printed scoring values change. Upgraded, Special, Golden and Overstrained Tiles cannot be upgraded.":
            "To Upgrade, spend 1 Action. Choose a reachable, placed basic Core Tile, pay the cost printed on its upgraded side and flip it in place. Its upgraded name, effect and scoring values now apply. Upgraded, Special, Golden and Overstrained Tiles cannot be upgraded.",
        "Spend 1 Action to use a reachable, non-Overstrained Tile with an Activated Effect. Passive effects require no Action. A Tile may be activated more than once if every activation is otherwise legal.":
            "To Activate, spend 1 Action to use a reachable, non-Overstrained Tile with an Activated Effect. A Tile may be activated more than once when each activation is legal. Passive effects never use the Activate Action.",
        "Your Steward Token follows the Tile you most recently changed. After the first placement, Place, Upgrade or Activate a Tile, move the Token onto that Tile; the next reachability check starts there.":
            "After your first placement, only Place, Upgrade or Activate moves your Steward Token. Move it to the Tile you just placed, upgraded or activated; begin your next reachability check there. Interact does not move the Token.",
        "Resource Tiles use Production as their Activated Effect. Add the printed resources to the Warehouse, up to its limit.":
            "Production is an Activated Effect on Resource Tiles. Spend an Activate Action, then add the Tile’s printed resources to the Warehouse, up to its limit.",
        "The first time an adjacent group of matching Resource Tiles is activated each round, the chosen Tile and every immediately adjacent, non-Overstrained matching Resource Tile produce for the single Activate Action.":
            "To resolve Linked Production, check the chosen Resource Tile and its immediately adjacent copies from the same Tile stack. If none has participated in Linked Production this round, the chosen Tile and every immediately adjacent, non-Overstrained copy produce for that single Activate Action. Use each producer’s current side. Linked Production does not continue outward through a chain.",
        "Further activations involving that group during the same round produce only from the chosen Tile. A separate disconnected matching group receives its own first Linked Production. Every group refreshes at the start of a round.":
            "Treat every Tile that produced as having participated this round. If the chosen Tile or any immediately adjacent matching Tile has already participated, only the chosen Tile produces. A disconnected matching set may still use its own Linked Production. All participation refreshes at the start of the next round.",
        "Workshops: basic and upgraded":
            "Workshops and The Makers Conclave",
        "Workshops: Once per round, when upgrading an adjacent Core Tile, reduce its cost by 1 resource. Its upgraded side, The Makers Conclave, changes the range: once per round, when upgrading a reachable Core Tile, reduce its cost by up to 2 resources.":
            "Workshops and The Makers Conclave are Passive Effects; neither uses an Activate Action. Once per round, Workshops reduces the cost to upgrade an adjacent Core Tile by 1 resource. Its upgraded side, The Makers Conclave, instead reduces the cost to upgrade any reachable Core Tile by up to 2 resources.",
        "Market Stalls: converting Goods":
            "Market Stalls and The Seldes",
        "Market Stalls: Once per round, when paying a cost, you may spend 1 Goods as 1 resource of any type. Its upgraded side, The Seldes, lets you spend 1 Goods as up to 2 resources of one type. The printed Goods requirement is never replaced by this conversion.":
            "Market Stalls and The Seldes are Passive Effects; neither uses an Activate Action. Once per round while paying a cost, Market Stalls lets you spend 1 Goods as 1 resource of any type. The Seldes instead lets you spend 1 Goods as up to 2 resources of one type. This conversion never replaces Goods printed in the cost.",
        "Every Tile placement passes the same checks: choose an available Tile, choose an empty reachable hex, check its terrain, adjacency and printed placement rule, pay the listed cost, then place the Tile and resolve any placement effect. The Vale contains Grasslands, Woodland, Arable Land, Mountains, Heaths, Ruins and River/Water terrain. Use the map terrain legend below to identify each region before applying a placement rule. Read the Tile’s printed requirement exactly: Mine Shaft goes on Mountains, Farmstead on Arable Land, Salvage Camp on Ruins and Bridge on Water/River. Ruins are special terrain: they provide a Stone discount. Any Land means an empty non-River hex unless the Tile’s own rule says otherwise.":
            "Every placement follows the same sequence: choose an available Tile and a reachable empty hex; check terrain, adjacency and the Tile’s printed placement rule; pay the cost; then place the Tile and resolve its placement effect. The Vale contains Grasslands, Woodland, Arable Land, Mountains, Heaths, Ruins and Water/River terrain. Use the map legend below to identify each region. Read printed requirements exactly: Mine Tunnel goes on Mountains, Farmstead on Arable Land, Salvage Camp on Ruins and Bridge on Water/River. Shelter in the Ruins: when placing or upgrading a Tile whose entire footprint lies on Ruins, ignore all Stone in that cost. If any part of the footprint lies outside Ruins, pay the full printed cost. Any Land means an empty non-Water/River hex unless the Tile says otherwise.",
        "Two Tiles are adjacent when their hexes share a side. Placement markers such as ADJ. HOUSING, ADJ. TRAVEL and ADJ. RIVER refer to this side-sharing relationship. Adjacency is not the same as reachability: a Tile may be reachable through the connected settlement network without being adjacent. All side-adjacent, non-Overstrained Tiles form the connected settlement network.":
            "Tiles are adjacent when their hexes share a side; a corner touch is not adjacency. Placement markers such as ADJ. HOUSING, ADJ. TRAVEL and ADJ. RIVER use this side-sharing rule. Adjacency is different from reachability: a Tile can be reachable through the connected settlement network without being adjacent. All side-adjacent, non-Overstrained Tiles form that network.",
        "A Tile or empty hex is reachable when the acting Steward can reach it through the shared connected settlement network, or when a card, Tile or Steward rule makes it reachable. All Stewards may use the shared network. Your Steward Token marks your current position in that network. After your first placement, whenever you Place, Upgrade or Activate, move the Token onto the Tile you just placed, upgraded or activated; your next reachability check starts from that Tile.":
            "A Tile or empty hex is reachable when the acting Steward can reach it through the shared connected settlement network, or when a card, Tile or Steward rule says it is reachable. Every Steward may use the shared network. Your Steward Token marks where your next reachability check begins. After your first placement, move it whenever you Place, Upgrade or Activate; Interact does not move it.",
        "Rivers are barriers. Tiles on opposite banks do not connect unless a Bridge or another rule allows the crossing. A Tile across an unbridged River is not reachable through normal connection.":
            "Water/River terrain is a barrier. Tiles on opposite banks do not connect unless a Bridge or another rule allows the crossing. A Tile across unbridged Water/River terrain is not reachable through the normal network.",
        "Every Encounter is read from top to bottom. The title and italic flavour text identify its story; the rules appear in the printed sections below. Specific card text overrides a general rule only for the situation it names.":
            "Read every Encounter from top to bottom. The title and italic flavour text tell its story; rules appear in the sections below. A specific card rule overrides a general rule only for the situation it names.",
        "Flavour text is narrative only. The separate italic line inside a Boon or Burden Season row is rules text; on a Golden Boon, the italic line at the base is rules text.":
            "Flavour text is narrative only. The separate italic line inside a Boon or Burden Season row is a rule. On a Golden Boon, the italic line at the base is also a rule.",
        "Boons help the settlement and are resolved from the row matching the current Season. Resolve the effect immediately unless the italic line says to keep the card face up. A kept Boon remains active until it is used or its printed duration ends, then it is discarded. For example, A Little Time adds 1 timer to an active Arrival in Season I, up to 3; its Season II and III rows add up to 2 or 3 timers as printed. Only one Boon may modify a single Action unless a rule says otherwise.":
            "Boons help the settlement. Read the row for the current Season and resolve it immediately unless the italic line tells you to keep the card face up. A kept Boon remains active until used or until its printed duration ends, then discard it. Using a face-up Boon costs no Action unless its text says otherwise. For example, A Little Time adds 1 timer to an active Arrival in Season I, up to 3; its later Season rows add up to 2 or 3 timers as printed. Only one Boon may modify a single Action unless a rule says otherwise.",
        "Burdens create ongoing pressure. Apply the current-Season effect when the card is revealed, then place it face up as an active Burden. At the start of Seasons II and III, every unresolved Burden reapplies using the new Season effect.":
            "Burdens create ongoing pressure. When one is revealed, apply the current-Season effect and leave the card face up as an active Burden. At the start of Seasons II and III, reapply every unresolved Burden using its new Season effect.",
        "Resolve a Burden by paying and performing its printed resolution, normally with an Interact Action, or by using an effect that explicitly resolves it. For example, Old Wounds Reopen lets you choose Social/Wellbeing Tiles and either pay Herbs or take Strain; its italic line gives the Action and payment needed to discard it. If a Burden offers payment or Strain, choose the branch before drawing a Target Card; draw only if the Strain branch is chosen. Resolution prevents later reapplication but does not remove Strain already placed.":
            "Resolve a Burden by paying and performing its printed resolution, usually with an Interact Action, or by using an effect that explicitly resolves it. For example, Old Wounds Reopen asks you to choose Social/Wellbeing Tiles, then pay Herbs or take Strain; its italic line states the Action and payment needed to discard it. If a Burden offers payment or Strain, choose the branch before drawing a Target Card and draw only for the Strain branch. Resolving a Burden prevents future reapplication but does not remove Strain already placed.",
        "Arrivals are timed objectives that unlock Special Tiles. Place a revealed Arrival face up with 3 timer tokens, one in each circle. For example, Quiet Quest requires 4 Goods and 2 Herbs and rewards the Adventurers’ Guild. To complete it, fulfil its Requirements and spend 1 Interact Action before the final timer is removed. Remove its timers, discard the Arrival and set its unlocked Special Tile beside the board; completing an Arrival unlocks its Tile but does not place it immediately. Timers, expiry and Inn activations are shown in the Arrival examples below. An Inn adds 1 timer to an active Arrival when activated; an upgraded Dawn Break Inn adds up to 2, and no Arrival can hold more than 3 timers.":
            "Arrivals are timed objectives that unlock Special Tiles. When one is revealed, place it face up with 3 timer tokens, one in each circle. Quiet Quest, for example, requires 4 Goods and 2 Herbs and unlocks the Adventurers’ Guild. To complete an Arrival, meet its Requirements and spend 1 Interact Action before its final timer is removed. Remove its timers, discard it and set the unlocked Special Tile beside the board. Any player may later place that Tile with a normal Place Action, following its printed placement rule and paying 0 resources. An activated Inn adds 1 timer to an active Arrival; Dawn Break Inn adds up to 2. An Arrival can never hold more than 3 timers.",
        "A Tile can hold up to 3 Strain. At 1 or 2 it functions normally unless a rule says otherwise. If an effect tries to add Strain to a Tile already at 3, add nothing.":
            "A Tile can hold at most 3 Strain. At 1 or 2 Strain, it functions normally unless a rule says otherwise. If an effect would add Strain to a Tile already at 3, add nothing.",
        "Supported prevents Strain before checking whether a Tile becomes Overstrained. A Tile either has Supported or does not; multiple instances do not stack. Some Core Tiles provide passive Supported: Common Land protects 1 adjacent Housing Tile, while its upgraded side, The Pleasance, protects up to 3. The source must not be Overstrained.":
            "Supported prevents 1 Strain before you check whether a Tile becomes Overstrained. A Tile either has Supported or does not; passive and single-use Supported never stack. Common Land gives passive Supported to 1 adjacent Housing Tile; The Pleasance gives it to up to 3. Choose the eligible Housing Tiles when the effect becomes active. Keep those choices until the eligible set changes or the source becomes Overstrained, then choose again when required. Passive Supported cannot be reassigned in response to incoming Strain. Supported granted by placement, activation or another effect is single-use unless the rule says it is passive.",
        "A Tile becomes Overstrained when it crosses from fewer than 3 Strain to 3. It remains on the map but cannot activate, upgrade, provide passive effects, connect the network, contribute Population or Renown, or satisfy positive requirements.":
            "A Tile becomes Overstrained when it crosses from fewer than 3 Strain to 3. Leave it on the map, but it cannot Activate, Upgrade, provide passive effects, connect the network, contribute Population or Renown, or satisfy positive requirements.",
        "As soon as it falls below 3 Strain, it works normally again. If it later reaches 3, that is a new Overstrain event.":
            "As soon as the Tile falls below 3 Strain, it works normally again. If it later returns to 3, that is a new Overstrain event.",
        "Overstrain spreads are resolved after the effect that caused them—not at the end of a round or Season.":
            "Resolve Overstrain spreads only after finishing the effect that caused them. Do not wait until the end of a round or Season.",
        "10.1 Finish the entire original effect, including all printed Strain placements and Strain Cascades.":
            "10.1  Finish the original effect, including every printed Strain placement and Strain Cascade.",
        "10.2 Identify every Tile that crossed from fewer than 3 Strain to 3.":
            "10.2  Identify every Tile that crossed from fewer than 3 Strain to 3.",
        "10.3 If several Tiles became Overstrained, the players choose their spread order.":
            "10.3  If several Tiles became Overstrained, choose the order in which they spread.",
        "10.4 For each spread, draw a fresh Target Card and apply its filters only to eligible adjacent placed Tiles below 3 Strain.":
            "10.4  For each spread, draw a fresh Target Card. Apply its filters only to eligible, adjacent placed Tiles below 3 Strain.",
        "10.5 Apply 1 Strain to the selected Tile, resolving Supported and other prevention.":
            "10.5  Apply 1 Strain to the selected Tile, resolving Supported and other prevention first.",
        "10.6 If the target becomes Overstrained, it also spreads. Continue until no new spread remains.":
            "10.6  If that Tile becomes Overstrained, it also spreads. Continue until no new spread remains.",
        "A card-specific Strain Cascade is part of the original effect. Complete every printed placement and Cascade before beginning the universal Overstrain chain they created.":
            "A printed Strain Cascade belongs to the original effect. Resolve every printed placement and Cascade before starting any universal Overstrain chain they create.",
        "Every Steward has a Starting Terrain, a Power usable once each Season and an Objective worth +15 Renown at final scoring. Powers do not require a Tile activation and refresh at the start of a Season.":
            "Each Steward has a Starting Terrain, a Power usable once per Season and an Objective worth +15 Renown at final scoring. Powers do not require a Tile activation and refresh at the start of each Season.",
        "Only non-Overstrained Tiles count toward positive Steward Objectives unless a rule says otherwise. Use the exact wording on each Steward Player Aid.":
            "Only non-Overstrained Tiles count toward positive Steward Objectives unless a rule says otherwise. Follow the exact wording on each Steward Aid.",
        "Neighbourly Support: Housing Tile Passive Effect":
            "Neighbourly Support",
        "At the end of Seasons I and II, each Housing cluster gains 1 single-use Supported for every 3 non-Overstrained Housing Tiles in that cluster. Place each Supported on a different non-Overstrained Housing Tile within that cluster. Ignore Overstrained Tiles when determining clusters and cluster size.":
            "At the end of Seasons I and II, each Housing cluster gains 1 single-use Supported for every complete group of 3 non-Overstrained Housing Tiles. Place each one on a different eligible Housing Tile in that cluster. Ignore Overstrained Tiles when tracing the cluster and counting its size.",
        "After Rounds 4 and 8, finish Arrival expiry and discard unused face-up Boons whose duration ends with the Season. Resolve Neighbourly Support before unresolved Burdens reapply for the new Season. Then refresh Steward Powers and seed Encounters for the next Season. There is no automatic Season-end Strain spread.":
            "After Rounds 4 and 8, resolve these steps in order: (1) finish every Arrival expiry; (2) discard unused face-up Boons whose duration ends with the Season; (3) resolve Neighbourly Support; (4) reapply every unresolved Burden using the new Season row; (5) refresh Steward Powers; and (6) seed Encounters for the next Season. There is no automatic Season-end Strain spread.",
        "After Round 12 and all required pending effects, calculate one shared score.":
            "After Round 12 and all required pending effects, calculate one shared final score. There is no pass/fail threshold or individual winner.",
        "The Ledger records achievements across games. Each named Ledger Entry counts once toward Golden unlocks; extra player-count boxes are prestige records only. Entries are never spent and do not make one player the winner.":
            "The Steward’s Ledger tracks achievements across games. Before your group’s first game, choose a name and an unused Roman numeral (I–VI) in the Group Register. Use that numeral for every Entry mark and Game Record. After scoring, mark NAMED for each newly completed Entry and record the shared final score and those Entries. Each NAMED mark counts once toward Golden milestones; player-count prestige marks are records only. Completed Entries are never spent or reset. An AFTER gate is met only if the group begins the game with the shown number of Named Entries.",
        "A Vow must be declared before setup and completed under its printed condition. Record completed Entries after final scoring.":
            "Declare a Vow before setup. A group may attempt only one Vow per game and completes it only under its printed condition.",
        "Use each Golden component’s printed rule. Every Golden Tile has a conditional +5 Renown scoring hook. The Golden Boon format is shown in the Encounter card guide on page 14: read its single Effect block, then follow the italic instruction at the base.":
            "Use each Golden component’s printed rule. Every Golden Tile has a conditional +5 Renown scoring hook.",
        "FOR DEFINITIONS, TIMING AND EDGE CASES, USE THE COMPANION RULES REFERENCE v6.5.":
            "USE THIS GLOSSARY AND THE RELEVANT RULEBOOK SECTION FOR DEFINITIONS, TIMING AND EDGE CASES.",
    }
    for old, new in paragraph_replacements.items():
        replace_paragraph(document, old, new)

    cell_replacements = {
        "Each player seeds one hidden Encounter on top of the deck, one into the middle and one on the bottom.":
            "Each player secretly seeds one Encounter on top of the deck, one into the middle and one on the bottom.",
        "Reveal Encounters → resolve them → each player takes a 4-Action turn → complete upkeep.":
            "Reveal and resolve Encounters → choose player order → each player takes one turn with up to 4 Actions → complete upkeep.",
        "After four rounds, discard expiring Boons; at the end of Seasons I and II, resolve Neighbourly Support before beginning the next Season. After Season III, score.":
            "After Rounds 4 and 8, resolve the Season transition in order. After Round 12, calculate the final score.",
        "Reveal standard Encounter Cards equal to the number of players. Resolve mandatory pending effects before continuing.":
            "Reveal until the number of standard Encounter Cards revealed equals the player count. Resolve every mandatory effect before Player Turns begin.",
        "Every player takes one turn with 4 Actions. The group chooses the order in which players act; each acts exactly once.":
            "Choose the player order. Each player then takes exactly one turn with up to 4 Actions.",
        "Remove 1 timer from every active Arrival. Expire those that reach 0. Discard effects whose duration ends, refresh round-based effects and advance the Round marker.":
            "Remove 1 timer from every active Arrival and resolve each expiry. Discard effects whose duration ends, refresh used round-based effects, then advance the Round marker.",
        "Add an available Tile to a legal, reachable empty hex and pay its cost.":
            "Spend 1 Action to place an available Tile on a legal, reachable empty hex and pay any required cost.",
        "Pay the upgrade cost and flip a reachable basic Core Tile to its upgraded side.":
            "Spend 1 Action, pay the cost printed on the upgraded side and flip a reachable basic Core Tile.",
        "Use a reachable, non-Overstrained Tile’s Activated Effect, including Production.":
            "Spend 1 Action to use a reachable, non-Overstrained Tile’s Activated Effect, including Production.",
        "Complete an Arrival or perform a printed interaction with an active Burden.":
            "Spend 1 Action to complete an active Arrival or perform an active Burden’s printed resolution.",
        "The Arrival fails. Discard it, record one failed Arrival and apply 1 Strain to the Tile selected by the Target Deck.":
            "The Arrival fails. Discard it and record one failed Arrival. Build the eligible pool from all placed Tiles below 3 Strain, then draw a Target Card to select one and apply 1 Strain. If no such Tile exists, place no Strain.",
        "Passive condition":
            "Passive Supported",
        "Refreshes each round while the condition remains true. Prevents the first Strain applied to that Tile that round.":
            "Applies while its printed condition remains true. Prevents the first Strain applied to that Tile each round.",
        "Placement, activation or other grant":
            "Single-use Supported",
        "Single-use. Prevents 1 Strain, then is spent. It returns only when another effect grants it.":
            "Prevents 1 Strain, then is spent. It returns only when another effect grants it.",
        "Place Travel for 0 Actions; it gains Supported beside Travel or on River/Water. Pay cost.":
            "Place Travel for 0 Actions; it gains Supported beside Travel or on Water/River. Pay cost.",
        "Build on 3+ non-Grasslands, non-River terrains.":
            "Build on 3+ non-Grasslands, non-Water/River terrains.",
        "Reveal and resolve Encounters":
            "Reveal and resolve Encounters",
        "Every player takes one 4-Action turn":
            "Choose order; each player takes one turn with up to 4 Actions",
        "Arrival timers, expiry, cleanup, advance round":
            "Remove Arrival timers, resolve expiry, clean up, advance round",
        "Add a legal Tile":
            "Place a legal Tile",
        "Flip a basic Core Tile":
            "Flip a reachable basic Core Tile",
        "Use an Activated Effect":
            "Use a reachable Activated Effect",
        "Complete Arrival / use Burden interaction":
            "Complete Arrival / resolve active Burden",
        "Draw Target Card within eligible pool":
            "Build eligible pool, then draw a Target Card",
        "Resolve Supported and other prevention":
            "Resolve Supported and other prevention first",
        "At 3 Strain, Tile becomes Overstrained":
            "On crossing to 3 Strain, Tile becomes Overstrained",
        "After original effect, each new Overstrained Tile spreads":
            "After the original effect, each newly Overstrained Tile spreads",
        "Available to the acting Steward through the network or a specific rule.":
            "A Tile or empty hex the acting Steward can reach through the connected network or a specific rule.",
        "A Tile unlocked by completing an Arrival and placed later with a normal Place Action.":
            "A Tile unlocked by an Arrival and later placed with a normal Place Action for 0 resources, following its printed placement rule.",
    }
    for old, new in cell_replacements.items():
        if old != new:
            replace_cell(document, old, new)

    setup_note = cell_paragraph_by_text(
        document,
        "FIRST PLACEMENT On your first turn, your first Action must place any legal Tile on the hex occupied by your Steward Token. It need not be a Resource Tile and need not connect to another Tile.",
    )
    set_labeled_paragraph(
        setup_note,
        "FIRST PLACEMENT",
        "On your first turn, your first Action must Place any legal Tile on the hex occupied by your Steward Token. The Tile need not be a Resource Tile or connect to another Tile.",
    )

    seeding_note = cell_paragraph_by_text(
        document,
        "WHY SEEDING MATTERS A Boon placed on top may help immediately. A Burden pushed deeper may arrive when the Vale is stronger. An Arrival delayed too long may crowd the final Season. You control timing, not certainty.",
    )
    set_labeled_paragraph(
        seeding_note,
        "TWO TIMING DECISIONS",
        "First choose which 3 cards enter the deck this Season; the rest stay in your hand. This is your clearest control over the Season row. Then assign top, middle and bottom. Top usually appears sooner, while bottom is most likely to cross a Season boundary. Position guides timing; it never fixes an exact round.",
    )
    seeding_table = table_containing(document, "TWO TIMING DECISIONS")
    add_seeding_decision_table(document, seeding_table)
    remove_paragraph_with_relationship(document, "rId31")

    golden_note = cell_paragraph_by_text(
        document,
        "GOLDEN BOONS A revealed Golden Boon resolves for free and does not count toward the standard reveal total. Continue revealing until the required number of standard Encounters has been revealed.",
    )
    set_labeled_paragraph(
        golden_note,
        "GOLDEN CONTENT",
        "Leave Golden Tiles and Golden Boons in the box until the Steward’s Ledger unlocks them. If your group includes unlocked Golden content, follow Section 16.",
    )

    chain_note = cell_paragraph_by_text(
        document,
        "CHAIN EXAMPLE Cottage and adjacent Path each have 2 Strain. An effect adds 1 to Cottage. After the effect finishes, Cottage spreads to the eligible adjacent Tile selected by a Target Card. If that Tile is Path, Path reaches 3 and later makes its own spread.",
    )
    set_labeled_paragraph(
        chain_note,
        "CHAIN EXAMPLE",
        "Cottage and an adjacent Path each have 2 Strain. An effect adds 1 Strain to Cottage. Finish that effect, then draw a Target Card for Cottage’s spread. If it selects Path, Path reaches 3 and immediately joins the same Overstrain chain with its own spread.",
    )

    final_note = cell_paragraph_by_text(
        document,
        "FINAL SCORE Add Population and all positive Renown, then subtract every penalty. Record the final total and any completed Steward’s Ledger Entries.",
    )
    set_labeled_paragraph(
        final_note,
        "FINAL SCORE",
        "Add Population and all positive Renown, then subtract every penalty. Record the shared final score and any newly completed Entries in the Steward’s Ledger.",
    )

    milestone_note = cell_paragraph_by_text(
        document,
        "CURRENT MILESTONES Golden content unlocks after 5, 12, 18, 25 and 32 completed named Ledger Entries. Previously earned milestones remain unlocked.",
    )
    set_labeled_paragraph(
        milestone_note,
        "CURRENT MILESTONES",
        "A group reaches Golden milestones at 5, 12, 18, 25 and 32 Named Entries. Each milestone unlocks 1 Golden Tile and 1 Golden Boon; unlocked rewards remain available in later games.",
    )

    conflict_note = cell_paragraph_by_text(
        document,
        "WHEN A RULE CONFLICTS Specific printed card or Tile text overrides the general rule only for the situation it names. Otherwise, follow this rulebook and resolve mandatory pending effects before play continues.",
    )
    set_labeled_paragraph(
        conflict_note,
        "WHEN RULES CONFLICT",
        "Specific printed card or Tile text overrides a general rule only for the situation it names. Otherwise, follow this rulebook and finish every mandatory pending effect before play continues.",
    )


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)
    apply_editorial_sweep(document)

    score_summary = table_containing(document, "YOUR SHARED SCORE")
    replace_run_text(
        score_summary.cell(0, 0).paragraphs[0],
        " − Unfulfilled Promises",
        "",
    )

    hidden_hands = table_containing(document, "For a first teaching game")
    replace_run_text(
        hidden_hands.cell(0, 0).paragraphs[0],
        " For a first teaching game, the group may agree to play with open hands.",
        "",
    )

    arrival_table = table_containing(document, "Still active after Round 12")
    set_paragraph_text(
        arrival_table.rows[-1].cells[1].paragraphs[0],
        "It remains active. It does not expire, place Strain or reduce the final score.",
    )
    arrival_note = table_containing(
        document, "A failed Arrival and an Unfulfilled Promise are different"
    )
    arrival_note_paragraph = arrival_note.cell(0, 0).paragraphs[0]
    set_paragraph_text(arrival_note_paragraph, "ROUND 12 ARRIVALS  ")
    arrival_note_paragraph.runs[0].bold = True
    arrival_note_paragraph.add_run(
        "An Arrival still active after Round 12 remains active. It does not expire, place Strain or reduce the final score."
    )

    final_score_table = table_containing(document, "7. Unfulfilled Promises")
    unfulfilled_row_index = next(
        index
        for index, row in enumerate(final_score_table.rows)
        if compact(row.cells[0].text) == "7. Unfulfilled Promises"
    )
    remove_row(final_score_table, unfulfilled_row_index)
    set_paragraph_text(final_score_table.rows[5].cells[1].paragraphs[0], "Subtract 10 for each.")
    set_paragraph_text(
        final_score_table.rows[6].cells[1].paragraphs[0],
        "Subtract 10 for each Arrival that expired during play.",
    )
    set_paragraph_text(final_score_table.rows[7].cells[0].paragraphs[0], "7. Strain")
    set_paragraph_text(
        final_score_table.rows[7].cells[1].paragraphs[0],
        "Subtract 10 for every Strain token on the map.",
    )

    penalty_tables = [
        table
        for table in document.tables
        if len(table.rows) == 2
        and any(
            compact(cell.text) == "Unfulfilled Promise"
            for cell in table.rows[0].cells
        )
    ]
    if len(penalty_tables) != 1:
        raise RuntimeError(
            f"Expected one quick-reference penalty table; found {len(penalty_tables)}"
        )
    penalty_table = penalty_tables[0]
    unfulfilled_column = next(
        index
        for index, cell in enumerate(penalty_table.rows[0].cells)
        if compact(cell.text) == "Unfulfilled Promise"
    )
    remove_equal_width_column(penalty_table, unfulfilled_column)
    for cell in penalty_table.rows[1].cells:
        set_paragraph_text(cell.paragraphs[0], "−10")

    glossary = table_containing(
        document, "An Arrival still active and incomplete after Round 12"
    )
    glossary_row_index = next(
        index
        for index, row in enumerate(glossary.rows)
        if compact(row.cells[0].text) == "Unfulfilled Promise"
    )
    remove_row(glossary, glossary_row_index)

    remove_heading_section(document, "17. First-game teaching order")
    quick_reference = paragraph_by_text(document, "18. Quick reference")
    set_paragraph_text(quick_reference, "17. Quick reference")
    remove_preceding_page_break(document, relationship_id="rId42")
    remove_preceding_page_break(document, paragraph_text="Glossary")

    contents = table_containing(document, "Quick reference and glossary")
    quick_reference_row = next(
        row
        for row in contents.rows
        if compact(row.cells[0].text) == "Quick reference and glossary"
    )
    set_paragraph_text(quick_reference_row.cells[1].paragraphs[0], "23–24")

    update_footer_versions(document)
    document.core_properties.title = "The Quiet Vale — Main Rulebook v6.6"
    document.core_properties.subject = (
        "Complete rules for 1–5 players with Encounter Card visual guide"
    )
    document.save(OUTPUT)
    replace_docx_members(OUTPUT, VISUAL_REPLACEMENTS)
    print(OUTPUT)


if __name__ == "__main__":
    main()
