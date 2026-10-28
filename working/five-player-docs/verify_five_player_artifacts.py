from __future__ import annotations

import re
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[2]
DOCX = ROOT / "output/docx/rulebook-v6_6/The_Quiet_Vale_Main_Rulebook_v6_6_Card_Guide.docx"
RULEBOOK_PDF = ROOT / "output/pdf/rulebook-v6_6/The_Quiet_Vale_Main_Rulebook_v6_6.pdf"
LEDGER_PDF = ROOT / "output/pdf/stewards-ledger-v6_5/The_Quiet_Vale_Stewards_Ledger_v6_5.pdf"
WORKBOOK = ROOT / "design/component-list/The_Quiet_Vale_Master_Component_List.xlsx"


def normalized(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def docx_text(path: Path) -> str:
    document = Document(path)
    parts: list[str] = []
    parts.extend(paragraph.text for paragraph in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    for section in document.sections:
        parts.extend(paragraph.text for paragraph in section.header.paragraphs)
        parts.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return normalized("\n".join(parts))


def pdf_text(path: Path, expected_pages: int) -> str:
    reader = PdfReader(path)
    assert len(reader.pages) == expected_pages, (
        f"{path.name}: expected {expected_pages} pages, found {len(reader.pages)}"
    )
    return normalized("\n".join(page.extract_text() or "" for page in reader.pages))


def assert_contains(text: str, phrases: list[str], label: str) -> None:
    missing = [phrase for phrase in phrases if phrase not in text]
    assert not missing, f"{label}: missing {missing}"


def workbook_rows(path: Path, sheet_name: str) -> list[list[object | None]]:
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"

    with zipfile.ZipFile(path) as archive:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in root.findall(f"{{{main_ns}}}si"):
                shared.append("".join(node.text or "" for node in item.iter(f"{{{main_ns}}}t")))

        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationship_id = None
        for sheet in workbook.findall(f".//{{{main_ns}}}sheet"):
            if sheet.attrib.get("name") == sheet_name:
                relationship_id = sheet.attrib[f"{{{rel_ns}}}id"]
                break
        assert relationship_id, f"Workbook sheet not found: {sheet_name}"

        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        target = None
        for relationship in relationships.findall(f"{{{package_rel_ns}}}Relationship"):
            if relationship.attrib.get("Id") == relationship_id:
                target = relationship.attrib["Target"]
                break
        assert target, f"Workbook relationship not found: {relationship_id}"
        worksheet_path = target.lstrip("/")
        if not worksheet_path.startswith("xl/"):
            worksheet_path = f"xl/{worksheet_path}"
        worksheet = ET.fromstring(archive.read(worksheet_path))

        parsed: dict[int, dict[int, object | None]] = {}
        for cell in worksheet.findall(f".//{{{main_ns}}}c"):
            reference = cell.attrib["r"]
            match = re.fullmatch(r"([A-Z]+)(\d+)", reference)
            assert match
            letters, row_number_text = match.groups()
            column_number = 0
            for letter in letters:
                column_number = column_number * 26 + ord(letter) - 64
            row_number = int(row_number_text)
            cell_type = cell.attrib.get("t")
            value_node = cell.find(f"{{{main_ns}}}v")
            inline_node = cell.find(f"{{{main_ns}}}is")
            value: object | None = None
            if cell_type == "s" and value_node is not None:
                value = shared[int(value_node.text or "0")]
            elif cell_type == "inlineStr" and inline_node is not None:
                value = "".join(node.text or "" for node in inline_node.iter(f"{{{main_ns}}}t"))
            elif value_node is not None and value_node.text is not None:
                raw = value_node.text
                try:
                    number = float(raw)
                    value = int(number) if number.is_integer() else number
                except ValueError:
                    value = raw
            parsed.setdefault(row_number, {})[column_number] = value

        width = max((max(row) for row in parsed.values() if row), default=0)
        return [
            [parsed.get(row_number, {}).get(column) for column in range(1, width + 1)]
            for row_number in range(1, max(parsed, default=0) + 1)
        ]


def verify_workbook() -> None:
    rows = workbook_rows(WORKBOOK, "Steward's Ledger")
    header_index = next(index for index, row in enumerate(rows) if row and row[0] == "Entry ID")
    header = rows[header_index]
    assert "5p Target" in header, "Workbook: 5p Target header missing"
    id_column = header.index("Entry ID")
    five_player_column = header.index("5p Target")
    boxes_column = header.index("Boxes")
    entries = [row for row in rows[header_index + 1 :] if len(row) > id_column and isinstance(row[id_column], str) and row[id_column].startswith("LE-")]
    assert [row[id_column] for row in entries] == [f"LE-{number:03d}" for number in range(1, 51)]

    expected_five_player_targets = {
        "LE-001": 380,
        "LE-002": 190,
        "LE-003": 120,
        "LE-004": 18,
        "LE-006": 32,
        "LE-011": 10,
        "LE-012": 8,
        "LE-016": 24,
        "LE-019": 10,
        "LE-022": 9,
        "LE-023": 4,
        "LE-024": 6,
        "LE-026": 20,
        "LE-027": 0,
        "LE-031": 13,
        "LE-036": 3,
        "LE-038": 85,
        "LE-039": 230,
        "LE-041": 290,
        "LE-042": 30,
        "LE-049": 6,
    }
    by_id = {row[id_column]: row for row in entries}
    for entry_id, expected in expected_five_player_targets.items():
        actual = by_id[entry_id][five_player_column]
        assert actual == expected, f"Workbook {entry_id}: expected 5p target {expected}, found {actual}"

    prestige_rows = [row for row in entries if len(row) > boxes_column and isinstance(row[boxes_column], str) and "1p" in row[boxes_column].lower()]
    assert prestige_rows, "Workbook: no player-count prestige rows found"
    assert all("5p" in str(row[boxes_column]).lower() for row in prestige_rows), "Workbook: a prestige row is missing 5p"

    with zipfile.ZipFile(WORKBOOK) as archive:
        worksheets = [name for name in archive.namelist() if name.startswith("xl/worksheets/") and name.endswith(".xml")]
        assert not any(b't="e"' in archive.read(name) for name in worksheets), "Workbook contains formula error cells"


def main() -> None:
    for path in (DOCX, RULEBOOK_PDF, LEDGER_PDF, WORKBOOK):
        assert path.is_file(), f"Missing artifact: {path}"

    docx = docx_text(DOCX)
    assert_contains(
        docx,
        [
            "1–5 players",
            "45–150 minutes",
            "choose any 5 and return the sixth Steward to the box",
            "20 of each type: 60 cards total",
            "deal 45 cards and place the remaining 15 in the deck",
            "30 Target Cards",
            "Fewer active Burdens than the player count",
            "TWO TIMING DECISIONS",
        ],
        "DOCX rulebook",
    )

    rulebook = pdf_text(RULEBOOK_PDF, 28)
    assert_contains(
        rulebook,
        [
            "1-5 players",
            "45-150 minutes",
            "choose any 5 and return the sixth Steward to the box",
            "20 of each type: 60 cards total",
            "deal 45 cards and place the remaining 15 in the deck",
            "30 Target Cards",
            "Fewer active Burdens than the player count",
            "15 / 10 / 5 / 0 / 0",
        ],
        "PDF rulebook",
    )

    ledger = pdf_text(LEDGER_PDF, 32)
    assert_contains(
        ledger,
        [
            "1P-5P rows",
            "140 / 200 / 280 / 320 / 380",
            "90 / 130 / 190 / 240 / 290",
            "110 / 85 / 75 / 40 / 30",
            "65 / 75 / 75 / 80 / 85",
            "The Small Storehouse",
            "3-5 players only",
            "5P 30",
        ],
        "Steward's Ledger PDF",
    )
    assert ledger.count("5P") >= 18, "Steward's Ledger PDF: fewer than 18 five-player prestige labels"

    stale = [
        "1–4 players",
        "1-4 players",
        "1P-4P rows",
        "3-4 players only",
        "24 Target Cards",
        "5 Boons, 5 Burdens and 5 Arrivals",
    ]
    for label, text in (("DOCX rulebook", docx), ("PDF rulebook", rulebook), ("Steward's Ledger PDF", ledger)):
        present = [phrase for phrase in stale if phrase in text]
        assert not present, f"{label}: stale phrases found {present}"

    verify_workbook()

    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    task_list = (ROOT / "AUGUST_TASK_LIST.md").read_text(encoding="utf-8")
    review = (ROOT / "PHYSICAL_PRODUCT_REVIEW_2026-08-12.md").read_text(encoding="utf-8")
    assert "current v6.6 Rulebook output" in readme
    assert "1P-5P rows" in task_list
    assert "Five-player documentation addendum" in review
    assert "at 1–5 players" in review

    print("Five-player artifact verification passed")
    print(f"DOCX: {DOCX}")
    print(f"Rulebook PDF: {RULEBOOK_PDF} (28 pages)")
    print(f"Ledger PDF: {LEDGER_PDF} (32 pages)")
    print(f"Workbook: {WORKBOOK} (50 entries, 5p targets verified)")


if __name__ == "__main__":
    try:
        main()
    except AssertionError as error:
        print(f"VERIFICATION FAILED: {error}", file=sys.stderr)
        raise SystemExit(1)
