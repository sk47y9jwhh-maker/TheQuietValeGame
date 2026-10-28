from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "output/docx/rulebook-v6_4/The_Quiet_Vale_Main_Rulebook_v6_4_Card_Guide.docx"
OUTPUT = ROOT / "output/docx/rulebook-v6_5/The_Quiet_Vale_Main_Rulebook_v6_5_Card_Guide.docx"
SETUP_VISUAL = ROOT / "working/rulebook-v6_4-card-visuals/standard-setup.png"
SETUP_RELATIONSHIP_ID = "rId29"


def compact(value: str) -> str:
    return " ".join(value.split())


def paragraph_matching(document: Document, text: str):
    wanted = compact(text)
    matches = [paragraph for paragraph in document.paragraphs if compact(paragraph.text) == wanted]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph {wanted!r}; found {len(matches)}")
    return matches[0]


def replace_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def replace_table_text(document: Document, old: str, new: str) -> None:
    replacements = 0
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if old in paragraph.text:
                        replace_paragraph_text(paragraph, paragraph.text.replace(old, new))
                        replacements += 1
    if replacements != 1:
        raise RuntimeError(f"Expected one table replacement for {old!r}; found {replacements}")


def replace_versions(document: Document) -> None:
    replacements = 0
    for section in document.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            for paragraph in footer.paragraphs:
                for run in paragraph.runs:
                    if "v6.4" in run.text or "v6.3" in run.text:
                        run.text = run.text.replace("v6.4", "v6.5").replace("v6.3", "v6.5")
                        replacements += 1
    if replacements == 0:
        raise RuntimeError("Footer version text was not found")


def add_five_player_warehouse_row(document: Document) -> None:
    warehouse_table = next(
        (
            table
            for table in document.tables
            if table.rows
            and len(table.columns) == 2
            and compact(table.cell(0, 0).text) == "Players"
            and compact(table.cell(0, 1).text) == "Start with"
        ),
        None,
    )
    if warehouse_table is None:
        raise RuntimeError("Starting Warehouse table was not found")
    if any(compact(row.cells[0].text) == "5" for row in warehouse_table.rows[1:]):
        return

    new_row = deepcopy(warehouse_table.rows[-1]._tr)
    warehouse_table._tbl.append(new_row)
    five_row = warehouse_table.rows[-1]
    replace_paragraph_text(five_row.cells[0].paragraphs[0], "5")
    replace_paragraph_text(five_row.cells[1].paragraphs[0], "0 resources")


def replace_setup_visual(docx_path: Path) -> None:
    document = Document(docx_path)
    relationship = document.part.rels.get(SETUP_RELATIONSHIP_ID)
    if relationship is None:
        raise RuntimeError(f"Setup visual relationship {SETUP_RELATIONSHIP_ID} was not found")
    target_name = str(relationship.target_part.partname).lstrip("/")
    temporary = docx_path.with_name(docx_path.stem + ".setup-image.tmp.docx")
    with ZipFile(docx_path, "r") as source, ZipFile(temporary, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = SETUP_VISUAL.read_bytes() if item.filename == target_name else source.read(item.filename)
            target.writestr(item, data)
    temporary.replace(docx_path)


def request_field_refresh(document: Document) -> None:
    settings = document.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not SETUP_VISUAL.exists():
        raise FileNotFoundError(SETUP_VISUAL)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)

    replace_paragraph_text(
        paragraph_matching(document, "1–4 players • 45–120 minutes • Cooperative"),
        "1–5 players  •  45–150 minutes  •  Cooperative",
    )
    replace_paragraph_text(
        paragraph_matching(
            document,
            "Each player chooses a different Steward, takes its Player Aid and Steward Token, and reads its Starting Terrain, once-per-Season Power and Steward Objective aloud. There are no Starting Benefits.",
        ),
        "Each player chooses a different Steward, takes its Steward Aid and Steward Token, and reads its Starting Terrain, once-per-Season Power and Steward Objective aloud. There are no Starting Benefits. In a 1-player game, choose 1 Steward and use the 1-player values wherever shown. In a 5-player game, choose 5 of the 6 Stewards; the unchosen Steward is not used.",
    )
    replace_paragraph_text(
        paragraph_matching(
            document,
            "2.4 Shuffle all 24 Target Cards and place the deck face down. Do not create a discard pile.",
        ),
        "2.4  Shuffle all 30 Target Cards and place the deck face down. Do not create a discard pile and do not reshuffle during play.",
    )
    replace_paragraph_text(
        paragraph_matching(
            document,
            "2.5 For each player, take 5 Boons, 5 Burdens and 5 Arrivals to form a balanced Encounter pool.",
        ),
        "2.5  Separate the standard Encounter Cards by type and shuffle each type. For each player, take 4 Boons, 4 Burdens and 4 Arrivals; combine all selected cards into one 12-card-per-player pool. In a 5-player game, this is 20 of each type, 60 cards total.",
    )
    replace_paragraph_text(
        paragraph_matching(document, "2.6 Shuffle that pool and deal 9 hidden Encounter Cards to each player."),
        "2.6  Shuffle the combined selected pool. Deal 9 hidden Encounter Cards to each player.",
    )
    replace_paragraph_text(
        paragraph_matching(document, "2.7 Place 3 standard Encounter Cards per player into the starting Encounter Deck."),
        "2.7  Without adding any other cards, place the remaining 3 selected standard Encounter Cards per player face down as the starting Encounter Deck. At 5 players, deal 45 cards and place the remaining 15 in the deck.",
    )
    replace_paragraph_text(
        paragraph_matching(
            document,
            "2.8 Keep the remaining standard cards face down as the reserve. Players may not inspect them unless an effect instructs them to do so.",
        ),
        "2.8  Return every unselected standard Encounter Card to the box as the reserve. Keep it separate from the starting deck; players may not inspect it unless an effect instructs them to do so.",
    )
    add_five_player_warehouse_row(document)

    component_replacements = {
        "The flat-top hex map on which the shared settlement grows.":
            "1 flat-top hex Game Map on which the shared settlement grows.",
        "Basic settlement Tiles with an upgraded reverse side.":
            "63 double-sided Core Tiles across 21 Tile types. Each has a basic side and an upgraded reverse side.",
        "Tiles unlocked by completing Arrivals.":
            "26 one-sided Special Tile components across 25 designs, unlocked by completing Arrivals.",
        "Boons, Burdens and Arrivals used to shape each Season.":
            "90 Encounter Cards: 30 Boons, 30 Burdens, 25 Arrivals and 5 Golden Boons.",
        "The 24-card deck that selects Tiles for incoming Strain.":
            "30 Target Cards that select Tiles for incoming Strain.",
        "Six Steward Player Aids and Steward Tokens.":
            "1 Main Rulebook, 6 Steward Aids and 6 matching Steward Tokens.",
        "The shared store of Wood, Stone, Metal, Food, Herbs and Goods.":
            "90 colour-coded resource pieces: 15 each of Wood, Stone, Metal, Food, Herbs and Goods.",
        "Strain, Supported, Arrival timers and Round marker.":
            "30 black markers and 30 bronze markers. Strain is black; Supported and the Round marker are bronze; Arrival timers may use either colour.",
        "A persistent record of achievements across games.":
            "1 Steward's Ledger booklet for tracking achievements, scores and Golden milestones across multiple games.",
        "No active Burdens.": "Fewer active Burdens than the player count.",
    }
    for old, new in component_replacements.items():
        replace_table_text(document, old, new)

    replace_table_text(
        document,
        "The online prototype highlights the current row; on a physical card, follow the matching Roman numeral.",
        "On a physical card, follow the row with the matching Roman numeral.",
    )

    companion = paragraph_matching(
        document,
        "FOR DEFINITIONS, TIMING AND EDGE CASES, USE THE COMPANION RULES REFERENCE v6.4.",
    )
    replace_paragraph_text(
        companion,
        "FOR DEFINITIONS, TIMING AND EDGE CASES, USE THE COMPANION RULES REFERENCE v6.5.",
    )

    replace_versions(document)
    request_field_refresh(document)
    document.core_properties.title = "The Quiet Vale — Main Rulebook v6.5"
    document.core_properties.subject = "Complete rules for 1–5 players with Encounter Card visual guide"
    document.save(OUTPUT)
    replace_setup_visual(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
